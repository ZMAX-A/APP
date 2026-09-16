from __future__ import annotations

from dataclasses import replace
from datetime import datetime
from pathlib import Path
from typing import cast
from unittest.mock import Mock

import pytest
from appium.webdriver.common.appiumby import AppiumBy
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from selenium.common.exceptions import InvalidSessionIdException

import yanjia_automation.config as config_module
import yanjia_automation.excel.workbook as workbook_module
from yanjia_automation.config import Settings, load_settings
from yanjia_automation.driver import is_recoverable_driver_error
from yanjia_automation.excel.locators import (
    LocatorFormatError,
    parse_locator,
    parse_locator_candidates,
)
from yanjia_automation.excel.models import ExcelCase, ExcelStep
from yanjia_automation.excel.results import ExcelResult, ExcelResultWriter
from yanjia_automation.excel.runner import ASSERTION_NAMES, ExcelCaseRunner
from yanjia_automation.excel.validation import validate_cases
from yanjia_automation.excel.variables import VariableResolver
from yanjia_automation.excel.workbook import ExcelCaseRepository


def test_repository_loads_enabled_case_and_ordered_steps(tmp_path: Path) -> None:
    workbook_path = tmp_path / "cases.xlsx"
    _create_workbook(workbook_path)

    cases = ExcelCaseRepository(workbook_path).load_cases()

    assert [case.case_id for case in cases] == ["TC-DEMO-001"]
    assert [step.order for step in cases[0].steps] == [1, 2]
    assert cases[0].tags == ("smoke", "readonly")


def test_repository_uses_eager_workbook_mode(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook_path = tmp_path / "cases.xlsx"
    _create_workbook(workbook_path)
    calls: list[dict[str, object]] = []
    original_load_workbook = workbook_module.load_workbook

    def tracked_load_workbook(
        filename: str | Path, *, read_only: bool = False, data_only: bool = False
    ) -> Workbook:
        calls.append({"read_only": read_only, "data_only": data_only})
        return original_load_workbook(
            filename, read_only=read_only, data_only=data_only
        )

    monkeypatch.setattr(workbook_module, "load_workbook", tracked_load_workbook)

    cases = ExcelCaseRepository(workbook_path).load_cases(
        case_patterns=("TC-DEMO-001",)
    )

    assert [case.case_id for case in cases] == ["TC-DEMO-001"]
    assert calls == [{"read_only": False, "data_only": True}]


def test_repository_migrates_targeted_single_sheet_android_login_case(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "single-sheet.xlsx"
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "自动化测试用例"
    sheet.append(
        (
            "用例ID",
            "模块",
            "测试场景",
            "测试点",
            "优先级",
            "前置条件",
            "操作步骤",
            "元素定位器",
            "操作类型",
            "输入数据",
            "数据类型",
            "期望结果",
            "验证点",
            "断言类型",
            "超时(秒)",
            "备注",
            "实际结果",
            "是否执行",
            "执行分组",
        )
    )
    sheet.append(
        (
            "TC-LOGIN-001",
            "账号登录",
            "登录失败-错误账号",
            "验证错误账号登录提示",
            "P0",
            "打开登录页面",
            "输入账号并提交",
            "#username, #password",
            "input,input,click",
            "invalid-account|${TEST_PASSWORD}",
            "string",
            "提示：登录失败，请重试！",
            "错误提示可见",
            "text_contains",
            5,
            "",
            "NOT_RUN",
            "是",
            "B",
        )
    )
    workbook.save(workbook_path)
    workbook.close()

    cases = ExcelCaseRepository(workbook_path).load_cases(case_patterns=("TC-LOGIN-001",))

    assert [case.case_id for case in cases] == ["TC-LOGIN-001"]
    assert [step.action for step in cases[0].steps] == [
        "restart_to_login",
        "input",
        "input",
        "click",
        "assert",
    ]
    assert cases[0].steps[1].locator == ("id=com.xiaofutech.yanjia_ai:id/login_username_et")
    assert cases[0].steps[-1].locator == ("id=com.xiaofutech.yanjia_ai:id/cover_prompt_cl")


def test_repository_migrates_remaining_android_login_cases(tmp_path: Path) -> None:
    workbook_path = tmp_path / "remaining-login-cases.xlsx"
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "自动化测试用例"
    sheet.append(
        (
            "用例ID",
            "模块",
            "测试场景",
            "测试点",
            "优先级",
            "前置条件",
            "输入数据",
            "期望结果",
            "超时(秒)",
            "是否执行",
        )
    )
    sheet.append(
        (
            "TC-LOGIN-004",
            "账号登录",
            "登录失败-密码为空",
            "验证密码必填",
            "P1",
            "打开登录页面",
            "fixture_user",
            "提示密码必填",
            3,
            "是",
        )
    )
    sheet.append(
        (
            "TC-LOGIN-005",
            "账号登录",
            "登录失败-未选择门店",
            "验证门店必选",
            "P1",
            "打开登录页面",
            "fixture_user|fixture_password",
            "提示选择门店",
            3,
            "是",
        )
    )
    sheet.append(
        (
            "TC-LOGIN-006",
            "账号登录",
            "查看用户协议",
            "查看用户协议",
            "P2",
            "打开登录页面",
            None,
            "进入用户协议界面",
            3,
            "是",
        )
    )
    workbook.save(workbook_path)
    workbook.close()

    cases = ExcelCaseRepository(workbook_path).load_cases(
        case_patterns=("TC-LOGIN-004", "TC-LOGIN-005", "TC-LOGIN-006")
    )
    steps_by_case = {case.case_id: case.steps for case in cases}

    assert [case.case_id for case in cases] == [
        "TC-LOGIN-004",
        "TC-LOGIN-005",
        "TC-LOGIN-006",
    ]
    assert [step.action for step in steps_by_case["TC-LOGIN-004"]] == [
        "restart_to_login",
        "input",
        "clear",
        "click",
        "assert",
    ]
    assert steps_by_case["TC-LOGIN-004"][1].input_value == "${TEST_USERNAME}"
    assert steps_by_case["TC-LOGIN-004"][-1].expected == ".LoginActivity"
    assert [step.action for step in steps_by_case["TC-LOGIN-005"]] == [
        "restart_to_login",
        "input",
        "input",
        "click",
        "assert",
        "assert",
        "assert",
    ]
    assert steps_by_case["TC-LOGIN-005"][4].locator == (
        "id=com.xiaofutech.yanjia_ai:id/login_store_rv"
    )
    assert steps_by_case["TC-LOGIN-005"][5].assertion == "element_count_gte"
    assert [step.action for step in steps_by_case["TC-LOGIN-006"]] == [
        "restart_to_login",
        "click",
        "assert",
        "assert",
    ]
    assert steps_by_case["TC-LOGIN-006"][1].input_value == "0.84,0.5"
    assert steps_by_case["TC-LOGIN-006"][2].expected == ".WebViewActivity"
    assert steps_by_case["TC-LOGIN-006"][3].expected == "用户协议"


def test_repository_migrates_home_p0_single_sheet_cases(tmp_path: Path) -> None:
    workbook_path = tmp_path / "home-p0-cases.xlsx"
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "自动化测试用例"
    sheet.append(
        (
            "用例ID",
            "模块",
            "测试场景",
            "测试点",
            "优先级",
            "前置条件",
            "输入数据",
            "期望结果",
            "超时(秒)",
            "是否执行",
        )
    )
    definitions = (
        ("TC-HOME-001", "首页", "首页模块展示"),
        ("TC-HOME-008", "首页跳转", "跳转顾客档案"),
        ("TC-HOME-009", "首页跳转", "跳转搜索"),
        ("TC-HOME-010", "首页跳转", "跳转案例库"),
        ("TC-HOME-011", "首页跳转", "跳转设置"),
        ("TC-HOME-012", "首页跳转", "跳转美际学院"),
        ("TC-CASE-001", "案例库", "进入案例管理"),
    )
    for case_id, module, scenario in definitions:
        sheet.append(
            (
                case_id,
                module,
                scenario,
                "验证页面跳转",
                "P0",
                "已登录成功",
                None,
                "目标页面可见",
                5,
                "是",
            )
        )
    workbook.save(workbook_path)
    workbook.close()

    cases = ExcelCaseRepository(workbook_path).load_cases(
        case_patterns=("TC-HOME-*", "TC-CASE-001")
    )
    steps_by_case = {case.case_id: case.steps for case in cases}

    assert [case.case_id for case in cases] == [item[0] for item in definitions]
    assert [step.action for step in steps_by_case["TC-HOME-001"]] == [
        "restart_to_home",
        "assert",
        "assert",
        "assert",
        "assert",
        "assert",
    ]
    assert steps_by_case["TC-HOME-001"][1].locator == (
        "id=com.xiaofutech.yanjia_ai:id/main_records_ll"
    )
    assert steps_by_case["TC-HOME-008"][-1].expected == ".CustomerRecordsActivity"
    assert steps_by_case["TC-HOME-009"][1].locator == (
        "id=com.xiaofutech.yanjia_ai:id/main_search_tv"
    )
    assert steps_by_case["TC-HOME-010"][-1].expected == ".CaseActivity"
    assert steps_by_case["TC-HOME-011"][-1].expected == ".SetActivity"
    assert steps_by_case["TC-HOME-012"][-1].expected == ".WebViewPCActivity"
    assert steps_by_case["TC-CASE-001"][-1].expected == ".CaseActivity"
    assert next(case for case in cases if case.case_id == "TC-CASE-001").tags == (
        "tablet",
        "readonly",
        "case-library",
    )


def test_repository_migrates_customer_and_detail_p0_single_sheet_cases(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "customer-p0-cases.xlsx"
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "自动化测试用例"
    sheet.append(
        (
            "用例ID",
            "模块",
            "测试场景",
            "测试点",
            "优先级",
            "前置条件",
            "输入数据",
            "期望结果",
            "超时(秒)",
            "是否执行",
        )
    )
    definitions = (
        ("TC-CUSTOMER-020", "顾客列表", "顾客卡片字段验证"),
        ("TC-CUSTOMER-021", "顾客列表", "进入顾客详情"),
        ("TC-DETAIL-001", "顾客详情", "编辑资料-姓名必填"),
        ("TC-DETAIL-011", "顾客详情", "编辑资料-手机号必填"),
    )
    for case_id, module, scenario in definitions:
        sheet.append(
            (
                case_id,
                module,
                scenario,
                "验证 Android 页面契约",
                "P0",
                "已登录并准备受控顾客数据",
                None,
                "页面满足预期",
                5,
                "是",
            )
        )
    workbook.save(workbook_path)
    workbook.close()

    cases = ExcelCaseRepository(workbook_path).load_cases(
        case_patterns=("TC-CUSTOMER-*", "TC-DETAIL-*")
    )
    steps_by_case = {case.case_id: case.steps for case in cases}

    assert [case.case_id for case in cases] == [item[0] for item in definitions]
    assert [step.action for step in steps_by_case["TC-CUSTOMER-020"]] == [
        "restart_to_home",
        "click",
        "assert",
        "assert",
        "assert",
        "assert",
        "assert",
        "assert",
        "assert",
        "assert",
    ]
    assert steps_by_case["TC-CUSTOMER-020"][3].assertion == "element_count_gte"
    assert steps_by_case["TC-CUSTOMER-020"][3].expected == "1"
    assert steps_by_case["TC-CUSTOMER-020"][4].locator == (
        "id=com.xiaofutech.yanjia_ai:id/a_records_head_ifv"
    )
    assert steps_by_case["TC-CUSTOMER-021"][3].locator == (
        "id=com.xiaofutech.yanjia_ai:id/a_records_cl"
    )
    assert steps_by_case["TC-CUSTOMER-021"][4].expected == ".CustomerDetailActivity"
    assert next(case for case in cases if case.case_id == "TC-CUSTOMER-020").tags == (
        "tablet",
        "readonly",
        "customer",
        "requires_seed",
    )

    name_steps = steps_by_case["TC-DETAIL-001"]
    assert name_steps[3].input_value == "${YANJIA_MUTATION_CUSTOMER_QUERY}"
    assert name_steps[5].assertion == "element_count_equals"
    assert name_steps[5].expected == "1"
    assert name_steps[9].locator == "id=com.xiaofutech.yanjia_ai:id/customer_edit_v"
    assert name_steps[10].action == "clear"
    assert name_steps[10].locator == (
        "id=com.xiaofutech.yanjia_ai:id/customer_edit_username_et"
    )
    assert name_steps[-2].locator == ("id=com.xiaofutech.yanjia_ai:id/customer_edit_save_tv")
    assert name_steps[-1].assertion == "page_source_contains"
    assert name_steps[-1].expected == "姓名不能为空"

    phone_steps = steps_by_case["TC-DETAIL-011"]
    assert phone_steps[10].locator == (
        "id=com.xiaofutech.yanjia_ai:id/customer_edit_unique_et"
    )
    assert phone_steps[-1].expected == "手机号不能为空"
    assert next(case for case in cases if case.case_id == "TC-DETAIL-011").tags == (
        "tablet",
        "mutating",
        "customer",
        "requires_seed",
        "serial",
    )


def test_repository_migrates_readonly_detail_remark_cases(tmp_path: Path) -> None:
    workbook_path = tmp_path / "detail-remark-cases.xlsx"
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "自动化测试用例"
    sheet.append(
        (
            "用例ID",
            "模块",
            "测试场景",
            "测试点",
            "优先级",
            "前置条件",
            "输入数据",
            "期望结果",
            "超时(秒)",
            "是否执行",
        )
    )
    sheet.append(
        (
            "TC-DETAIL-023",
            "顾客详情",
            "影像备注-备注空格",
            "验证空格备注无法提交",
            "P1",
            "受控顾客已有影像",
            None,
            "点击提交后备注记录不增加",
            10,
            "是",
        )
    )
    sheet.append(
        (
            "TC-DETAIL-036",
            "顾客详情",
            "咨询单-备注空格",
            "验证空格备注无法提交",
            "P1",
            "受控顾客已有咨询单",
            "咨询单特定| ",
            "点击提交后备注记录不增加",
            10,
            "是",
        )
    )
    workbook.save(workbook_path)
    workbook.close()

    cases = ExcelCaseRepository(workbook_path).load_cases(
        case_patterns=("TC-DETAIL-023", "TC-DETAIL-036")
    )
    by_id = {case.case_id: case for case in cases}

    image_steps = by_id["TC-DETAIL-023"].steps
    assert image_steps[6].locator == (
        "id=com.xiaofutech.yanjia_ai:id/a_records_detail_all_remark_ifv"
    )
    assert image_steps[9].input_value == "${SPACE}"
    assert image_steps[10].assertion == "element_count_unchanged_after_click"
    assert image_steps[10].expected == (
        "id=com.xiaofutech.yanjia_ai:id/a_records_remark_list_content_tv"
    )
    assert all(
        step.locator != "id=com.xiaofutech.yanjia_ai:id/records_remark_remark_commit_tv"
        or step.action == "assert"
        for step in image_steps
    )

    consultation_steps = by_id["TC-DETAIL-036"].steps
    assert consultation_steps[3].input_value == "咨询单特定"
    assert consultation_steps[9].locator == (
        "id=com.xiaofutech.yanjia_ai:id/a_consultation_result_remark_ifv"
    )
    assert consultation_steps[12].input_value == "${SPACE}"
    assert consultation_steps[13].assertion == "element_count_unchanged_after_click"
    assert consultation_steps[13].expected == (
        "id=com.xiaofutech.yanjia_ai:id/a_records_remark_list_content_tv"
    )
    assert by_id["TC-DETAIL-023"].tags == (
        "tablet",
        "readonly",
        "customer",
        "requires_seed",
    )
    assert by_id["TC-DETAIL-036"].tags == by_id["TC-DETAIL-023"].tags


def test_variable_resolution_and_redaction() -> None:
    resolver = VariableResolver(
        {"NAME": "演示顾客", "PASSWORD": "secret-value"},
        sensitive_values={"secret-value"},
    )

    assert resolver.resolve("你好，${NAME}") == "你好，演示顾客"
    assert resolver.resolve("<EMPTY>") == ""
    assert resolver.redact("password=secret-value") == "password=***"


def test_element_disabled_assertion_checks_native_enabled_state(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    runner = object.__new__(ExcelCaseRunner)
    runner.variables = VariableResolver({}, sensitive_values=set())
    step = ExcelStep(
        source_row=2,
        case_id="TC-DETAIL-023",
        order=1,
        name="校验提交按钮禁用",
        action="assert",
        locator="id=com.xiaofutech.yanjia_ai:id/records_remark_remark_commit_tv",
        input_value=None,
        assertion="element_disabled",
        expected=None,
        timeout=1.0,
        element_index=0,
        continue_on_failure=False,
        enabled=True,
        note=None,
    )
    element = type("DisabledElement", (), {"is_enabled": lambda self: False})()
    observed: dict[str, object] = {}
    monkeypatch.setattr(runner, "_element", lambda *_args, **_kwargs: element)
    monkeypatch.setattr(
        runner,
        "_wait_condition",
        lambda predicate, _step, message: observed.update(result=predicate(), message=message),
    )

    runner._apply_assertion(cast(ExcelCase, None), step, "element_disabled")

    assert "element_disabled" in ASSERTION_NAMES
    assert observed == {"result": True, "message": "元素仍处于启用状态"}


def test_element_count_unchanged_after_click_uses_touch_and_observes_no_new_record() -> None:
    runner = object.__new__(ExcelCaseRunner)
    runner.variables = VariableResolver({}, sensitive_values=set())
    runner.driver = Mock()
    runner.driver.current_activity = ".RecordsRemarkActivity"
    existing_record = Mock()
    runner.driver.find_elements.return_value = [existing_record]
    button = Mock()
    button.rect = {"x": 100, "y": 200, "width": 80, "height": 40}
    runner._element = Mock(return_value=button)
    step = ExcelStep(
        source_row=2,
        case_id="TC-DETAIL-036",
        order=1,
        name="点击提交并校验备注未新增",
        action="assert",
        locator="id=com.xiaofutech.yanjia_ai:id/records_remark_remark_commit_tv",
        input_value=None,
        assertion="element_count_unchanged_after_click",
        expected="id=com.xiaofutech.yanjia_ai:id/a_records_remark_list_content_tv",
        timeout=0.01,
        element_index=0,
        continue_on_failure=False,
        enabled=True,
        note=None,
    )

    runner._apply_assertion(
        cast(ExcelCase, None), step, "element_count_unchanged_after_click"
    )

    runner.driver.execute_script.assert_called_once_with(
        "mobile: clickGesture", {"x": 140, "y": 220}
    )
    assert "element_count_unchanged_after_click" in ASSERTION_NAMES


def test_element_count_unchanged_after_click_fails_when_record_is_added() -> None:
    runner = object.__new__(ExcelCaseRunner)
    runner.variables = VariableResolver({}, sensitive_values=set())
    runner.driver = Mock()
    runner.driver.current_activity = ".RecordsRemarkActivity"
    runner.driver.find_elements.side_effect = [[Mock()], [Mock(), Mock()]]
    button = Mock()
    button.rect = {"x": 0, "y": 0, "width": 20, "height": 20}
    runner._element = Mock(return_value=button)
    step = ExcelStep(
        source_row=2,
        case_id="TC-DETAIL-036",
        order=1,
        name="点击提交并校验备注未新增",
        action="assert",
        locator="id=com.xiaofutech.yanjia_ai:id/records_remark_remark_commit_tv",
        input_value=None,
        assertion="element_count_unchanged_after_click",
        expected="id=com.xiaofutech.yanjia_ai:id/a_records_remark_list_content_tv",
        timeout=1.0,
        element_index=0,
        continue_on_failure=False,
        enabled=True,
        note=None,
    )

    with pytest.raises(AssertionError, match="备注记录数量发生变化"):
        runner._apply_assertion(
            cast(ExcelCase, None), step, "element_count_unchanged_after_click"
        )


def test_variable_resolver_inherits_credentials_for_empty_test_aliases(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        config_module,
        "_DOTENV",
        {
            "YANJIA_USERNAME": "android-user",
            "YANJIA_PASSWORD": "android-password",
            "TEST_USERNAME": "",
            "TEST_PASSWORD": "",
        },
    )
    for name in (
        "YANJIA_USERNAME",
        "YANJIA_PASSWORD",
        "YANJIA_IGNORE_DOTENV",
        "TEST_USERNAME",
        "TEST_PASSWORD",
    ):
        monkeypatch.delenv(name, raising=False)

    settings = Settings(
        appium_server_url="http://127.0.0.1:4723",
        app_package="com.xiaofutech.yanjia_ai",
        app_activity=".activity.SplashActivity",
        udid="device",
        store_name=None,
        no_reset=True,
        run_seeded=False,
        allow_mutation=False,
        allow_destructive=False,
        capture_sensitive_artifacts=False,
        skip_device_initialization=True,
        skip_server_installation=True,
    )

    resolver = VariableResolver.from_settings(settings, run_id="test")

    assert resolver.resolve("${TEST_USERNAME}") == "android-user"
    assert resolver.resolve("${TEST_PASSWORD}") == "android-password"


def test_offline_variable_resolver_does_not_require_credentials(tmp_path: Path) -> None:
    settings = replace(load_settings(), project_root=tmp_path)

    resolver = VariableResolver.from_settings(
        settings,
        run_id="VALIDATION",
        require_credentials=False,
    )

    assert resolver.resolve("${RUN_ID}") == "VALIDATION"
    assert resolver.resolve("${RUN_TOKEN}") == "DATION"
    run_phone = resolver.resolve("${RUN_PHONE}")
    assert run_phone is not None and len(run_phone) == 11 and run_phone.isdigit()
    assert run_phone.startswith("13900")
    assert resolver.resolve("${INVALID_USERNAME}") == "__invalid_yanjia_user__"
    assert resolver.resolve("${INVALID_PASSWORD}") == "__invalid_yanjia_password__"
    assert resolver.resolve("${NON_EXISTENT_CASE_TAG}") == "__missing_case_tag_DATION__"
    assert resolver.resolve("${SPACE}") == " "


def test_locator_parser() -> None:
    assert parse_locator("id=com.example:id/button") == (
        AppiumBy.ID,
        "com.example:id/button",
    )
    assert parse_locator("accessibility_id=返回") == (AppiumBy.ACCESSIBILITY_ID, "返回")
    assert parse_locator_candidates("id=com.example:id/button || accessibility_id=保存") == (
        (AppiumBy.ID, "com.example:id/button"),
        (AppiumBy.ACCESSIBILITY_ID, "保存"),
    )
    with pytest.raises(LocatorFormatError):
        parse_locator_candidates("id=button || ")


def test_validator_reports_missing_variable_and_bad_action(tmp_path: Path) -> None:
    workbook_path = tmp_path / "cases.xlsx"
    _create_workbook(workbook_path)
    case = ExcelCaseRepository(workbook_path).load_cases()[0]
    bad_step = replace(
        case.steps[0],
        action="unsupported_action",
        expected="$" + "{MISSING_VALUE}",
    )
    report = validate_cases(
        [replace(case, steps=(bad_step,))],
        VariableResolver({}, sensitive_values=set()),
    )

    assert not report.is_valid
    assert any("不支持的操作类型" in issue.message for issue in report.errors)
    assert any("MISSING_VALUE" in issue.message for issue in report.errors)


def test_infrastructure_error_classifier_never_retries_assertions() -> None:
    assert is_recoverable_driver_error(InvalidSessionIdException("invalid session id"))
    assert not is_recoverable_driver_error(AssertionError("页面业务断言失败"))


def test_result_writer_updates_latest_result_and_history(tmp_path: Path) -> None:
    workbook_path = tmp_path / "cases.xlsx"
    _create_workbook(workbook_path)
    writer = ExcelResultWriter(
        source_path=workbook_path,
        output_path=None,
        run_id="run-001",
        allure_report_dir=tmp_path / "allure-report",
        enabled=True,
    )
    writer.prepare(("TC-DEMO-001",))
    started = datetime(2026, 7, 30, 10, 0, 0)
    finished = datetime(2026, 7, 30, 10, 0, 1)

    writer.record(
        ExcelResult(
            case_id="TC-DEMO-001",
            status="PASS",
            started_at=started,
            finished_at=finished,
            duration_seconds=1.25,
        )
    )

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        cases = workbook["自动化测试用例"]
        headers: dict[str, int] = {}
        for cell in cases[1]:
            if isinstance(cell.value, str) and isinstance(cell.column, int):
                headers[cell.value] = cell.column
        assert cases.cell(2, headers["实际结果"]).value == "PASS"
        assert cases.cell(2, headers["运行编号"]).value == "run-001"
        history = workbook["执行记录"]
        assert history.cell(2, 2).value == "TC-DEMO-001"
        assert history.cell(2, 5).value == "PASS"
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("status", "expected_fill", "expected_font"),
    (
        ("PASS", "FFC6EFCE", "FF006100"),
        ("FAIL", "FFFFC7CE", "FF9C0006"),
        ("ERROR", "FFFFC7CE", "FF9C0006"),
    ),
)
def test_result_writer_colors_actual_result_by_status(
    tmp_path: Path,
    status: str,
    expected_fill: str,
    expected_font: str,
) -> None:
    workbook_path = tmp_path / f"cases-{status.lower()}.xlsx"
    _create_workbook(workbook_path)
    writer = ExcelResultWriter(
        source_path=workbook_path,
        output_path=None,
        run_id=f"run-{status.lower()}",
        allure_report_dir=tmp_path / "allure-report",
        enabled=True,
    )
    writer.prepare(("TC-DEMO-001",))
    writer.record(
        ExcelResult(
            case_id="TC-DEMO-001",
            status=status,
            started_at=datetime(2026, 9, 4, 15, 0, 0),
            finished_at=datetime(2026, 9, 4, 15, 0, 1),
            duration_seconds=1.0,
        )
    )

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        cases = workbook["自动化测试用例"]
        headers = {
            str(cell.value): cell.column
            for cell in cases[1]
            if cell.value and isinstance(cell.column, int)
        }
        result = cases.cell(2, headers["实际结果"])
        assert result.value == status
        assert result.fill.fill_type == "solid"
        assert result.fill.fgColor.rgb == expected_fill
        assert result.font.bold is True
        assert result.font.color is not None
        assert result.font.color.rgb == expected_font
    finally:
        workbook.close()


def _create_workbook(path: Path) -> None:
    workbook = Workbook()
    cases = cast(Worksheet, workbook.active)
    cases.title = "自动化测试用例"
    cases.append(
        (
            "用例ID",
            "模块",
            "测试场景",
            "测试点",
            "优先级",
            "前置条件",
            "输入数据",
            "期望结果",
            "超时(秒)",
            "实际结果",
            "自动化状态",
            "标签",
            "是否执行",
        )
    )
    cases.append(
        (
            "TC-DEMO-001",
            "演示模块",
            "演示场景",
            "演示测试点",
            "P1",
            "已进入首页",
            None,
            "按钮可见",
            5,
            "NOT_RUN",
            "AUTOMATED",
            "smoke,readonly",
            "是",
        )
    )
    cases.append(
        (
            "TC-DEMO-002",
            "演示模块",
            "禁用场景",
            "禁用测试点",
            "P2",
            None,
            None,
            None,
            5,
            "NOT_RUN",
            "BACKLOG",
            "readonly",
            "否",
        )
    )

    steps = workbook.create_sheet("自动化执行步骤")
    steps.append(
        (
            "用例ID",
            "步骤序号",
            "步骤名称",
            "操作类型",
            "元素定位器",
            "输入数据",
            "断言类型",
            "期望值",
            "超时(秒)",
            "元素索引",
            "失败继续",
            "是否执行",
            "备注",
        )
    )
    steps.append(
        (
            "TC-DEMO-001",
            2,
            "断言按钮",
            "assert",
            "id=com.example:id/button",
            None,
            "element_visible",
            None,
            5,
            0,
            "否",
            "是",
            None,
        )
    )
    steps.append(
        (
            "TC-DEMO-001",
            1,
            "恢复首页",
            "restart_to_home",
            None,
            None,
            None,
            None,
            30,
            0,
            "否",
            "是",
            None,
        )
    )
    workbook.save(path)
    workbook.close()

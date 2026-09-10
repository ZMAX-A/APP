from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from yanjia_automation.excel.android_catalog import (
    SINGLE_SHEET_SOURCE,
    STEP_SHEET_SOURCE,
    android_case_ids,
    get_android_case_definition,
)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "test_case.xlsx"


@dataclass(frozen=True, slots=True)
class WorkbookCase:
    case_id: str
    module: str
    priority: str


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Report Android case coverage from the Excel case sheet."
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument(
        "--source",
        choices=(SINGLE_SHEET_SOURCE, STEP_SHEET_SOURCE),
        default=SINGLE_SHEET_SOURCE,
        help="Android execution adapter to measure.",
    )
    parser.add_argument(
        "--fail-on-p0-backlog",
        action="store_true",
        help="Return exit code 2 when an enabled P0 case is not covered.",
    )
    args = parser.parse_args(argv)

    workbook_path = args.workbook.resolve()
    if not workbook_path.is_file():
        parser.error(f"Workbook not found: {workbook_path}")

    total, enabled, duplicate_ids = _read_cases(workbook_path)
    supported_ids = android_case_ids(source=args.source)
    supported = [case for case in enabled if case.case_id in supported_ids]
    backlog = [case for case in enabled if case.case_id not in supported_ids]

    print("Android coverage report")
    print(f"workbook: {workbook_path}")
    print(f"source: {args.source}")
    print(
        f"cases: total={total}, enabled={len(enabled)}, "
        f"supported={len(supported)}, backlog={len(backlog)}"
    )
    _print_priority_summary(enabled, supported, backlog)

    if duplicate_ids:
        print("duplicate IDs: " + ", ".join(duplicate_ids))

    print("supported IDs: " + ", ".join(sorted({case.case_id for case in supported})))
    if backlog:
        print("backlog:")
        for case in sorted(backlog, key=lambda item: (item.priority, item.case_id)):
            print(
                f"  [{case.priority or '-'}] {case.module or '-'} {case.case_id} - "
                f"{_backlog_reason(case.case_id, args.source)}"
            )

    p0_backlog = sum(case.priority == "P0" for case in backlog)
    if duplicate_ids or (args.fail_on_p0_backlog and p0_backlog):
        return 2
    return 0


def _read_cases(workbook_path: Path) -> tuple[int, list[WorkbookCase], list[str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "自动化测试用例" not in workbook.sheetnames:
            raise ValueError("Workbook is missing sheet: 自动化测试用例")
        sheet = workbook["自动化测试用例"]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in sheet[1]
            if cell.value not in (None, "") and isinstance(cell.column, int)
        }
        required = {"用例ID", "模块", "优先级", "是否执行"}
        missing = sorted(required - headers.keys())
        if missing:
            raise ValueError("Workbook is missing headers: " + ", ".join(missing))

        all_ids: list[str] = []
        enabled: list[WorkbookCase] = []
        for row in sheet.iter_rows(min_row=2, values_only=False):
            case_id = _text(row[headers["用例ID"] - 1].value)
            if not case_id:
                continue
            all_ids.append(case_id)
            if _enabled(row[headers["是否执行"] - 1].value):
                enabled.append(
                    WorkbookCase(
                        case_id=case_id,
                        module=_text(row[headers["模块"] - 1].value),
                        priority=_text(row[headers["优先级"] - 1].value).upper(),
                    )
                )
        duplicates = sorted(case_id for case_id, count in Counter(all_ids).items() if count > 1)
        return len(all_ids), enabled, duplicates
    finally:
        workbook.close()


def _print_priority_summary(
    enabled: list[WorkbookCase],
    supported: list[WorkbookCase],
    backlog: list[WorkbookCase],
) -> None:
    priorities = sorted(
        set(case.priority for case in enabled),
        key=lambda value: (value not in {"P0", "P1", "P2", "P3"}, value),
    )
    for priority in priorities:
        total_count = sum(case.priority == priority for case in enabled)
        supported_count = sum(case.priority == priority for case in supported)
        backlog_count = sum(case.priority == priority for case in backlog)
        print(
            f"{priority or '-'}: total={total_count}, "
            f"supported={supported_count}, backlog={backlog_count}"
        )


def _backlog_reason(case_id: str, source: str) -> str:
    definition = get_android_case_definition(case_id)
    if definition is None:
        return "未登记到 Android 用例目录"
    if definition.blocked_reason:
        return definition.blocked_reason
    return f"已登记，但尚未接入 {source}"


def _text(value: object) -> str:
    return str(value).strip() if value not in (None, "") else ""


def _enabled(value: object) -> bool:
    return _text(value).lower() in {"是", "yes", "y", "true", "1", "on"}


if __name__ == "__main__":
    raise SystemExit(main())

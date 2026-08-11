from __future__ import annotations

import os
import shutil
import threading
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from yanjia_automation.excel.workbook import CASE_SHEET

RESULT_COLUMNS = (
    "是否执行",
    "实际结果",
    "最后执行时间",
    "执行耗时(秒)",
    "错误信息",
    "运行编号",
    "Allure报告目录",
)

HISTORY_SHEET = "执行记录"
HISTORY_HEADERS = (
    "运行编号",
    "用例ID",
    "开始时间",
    "结束时间",
    "执行结果",
    "执行耗时(秒)",
    "错误信息",
    "Allure报告目录",
)


@dataclass(frozen=True)
class ExcelResult:
    case_id: str
    status: str
    started_at: datetime
    finished_at: datetime
    duration_seconds: float
    error_message: str = ""


class ExcelResultWriter:
    def __init__(
        self,
        *,
        source_path: Path,
        output_path: Path | None,
        run_id: str,
        allure_report_dir: Path,
        enabled: bool,
    ) -> None:
        self.source_path = source_path.resolve()
        self.target_path = (output_path or source_path).resolve()
        self.run_id = run_id
        self.allure_report_dir = allure_report_dir.resolve()
        self.enabled = enabled
        self._lock = threading.Lock()
        self._prepared = False

    def prepare(self, selected_case_ids: tuple[str, ...]) -> None:
        if not self.enabled or self._prepared:
            return
        with self._lock:
            if self._prepared:
                return
            if self.target_path != self.source_path:
                self.target_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(self.source_path, self.target_path)

            backup_dir = self.source_path.parent / "reports" / "excel-backups"
            backup_dir.mkdir(parents=True, exist_ok=True)
            backup_path = backup_dir / f"{self.source_path.stem}_before_{self.run_id}.xlsx"
            shutil.copy2(self.target_path, backup_path)

            workbook = self._load_target()
            try:
                sheet = workbook[CASE_SHEET]
                headers = _ensure_columns(sheet, RESULT_COLUMNS)
                selected = set(selected_case_ids)
                for row in range(2, sheet.max_row + 1):
                    case_id = str(sheet.cell(row, headers["用例ID"]).value or "").strip()
                    if case_id not in selected:
                        continue
                    sheet.cell(row, headers["实际结果"], "NOT_RUN")
                    sheet.cell(row, headers["最后执行时间"], "")
                    sheet.cell(row, headers["执行耗时(秒)"], "")
                    sheet.cell(row, headers["错误信息"], "")
                    sheet.cell(row, headers["运行编号"], self.run_id)
                    sheet.cell(row, headers["Allure报告目录"], str(self.allure_report_dir))
                _ensure_history_sheet(workbook)
                self._atomic_save(workbook)
            finally:
                workbook.close()
            self._prepared = True

    def record(self, result: ExcelResult) -> None:
        if not self.enabled:
            return
        if not self._prepared:
            raise RuntimeError("Excel结果写入器尚未初始化")
        with self._lock:
            workbook = self._load_target()
            try:
                sheet = workbook[CASE_SHEET]
                headers = _ensure_columns(sheet, RESULT_COLUMNS)
                row = _find_case_row(sheet, headers, result.case_id)
                sheet.cell(row, headers["实际结果"], result.status)
                sheet.cell(
                    row,
                    headers["最后执行时间"],
                    result.finished_at.isoformat(timespec="seconds"),
                )
                sheet.cell(row, headers["执行耗时(秒)"], round(result.duration_seconds, 3))
                sheet.cell(row, headers["错误信息"], result.error_message[:10_000])
                sheet.cell(row, headers["运行编号"], self.run_id)
                sheet.cell(row, headers["Allure报告目录"], str(self.allure_report_dir))

                history = _ensure_history_sheet(workbook)
                history.append(
                    (
                        self.run_id,
                        result.case_id,
                        result.started_at.isoformat(timespec="seconds"),
                        result.finished_at.isoformat(timespec="seconds"),
                        result.status,
                        round(result.duration_seconds, 3),
                        result.error_message[:10_000],
                        str(self.allure_report_dir),
                    )
                )
                self._atomic_save(workbook)
            finally:
                workbook.close()

    def _load_target(self):
        try:
            return load_workbook(self.target_path)
        except PermissionError as error:
            raise RuntimeError(
                f"无法写入Excel：{self.target_path}。请关闭正在打开该文件的Excel窗口。"
            ) from error

    def _atomic_save(self, workbook) -> None:
        temporary = self.target_path.with_name(
            f".{self.target_path.stem}.{self.run_id}.tmp.xlsx"
        )
        try:
            workbook.save(temporary)
            os.replace(temporary, self.target_path)
        except PermissionError as error:
            raise RuntimeError(
                f"无法覆盖Excel：{self.target_path}。请关闭正在打开该文件的Excel窗口。"
            ) from error
        finally:
            temporary.unlink(missing_ok=True)


def _ensure_columns(sheet: Worksheet, names: tuple[str, ...]) -> dict[str, int]:
    headers = _header_map(sheet)
    style_source = sheet.cell(1, headers.get("实际结果", 1))
    for name in names:
        if name in headers:
            continue
        column = sheet.max_column + 1
        target = sheet.cell(1, column, name)
        target.font = Font(
            name=style_source.font.name,
            size=style_source.font.size,
            bold=True,
            color=style_source.font.color,
        )
        target.fill = PatternFill(
            fill_type=style_source.fill.fill_type,
            fgColor=style_source.fill.fgColor,
        )
        target.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[target.column_letter].width = 22
        headers[name] = column
    return headers


def _header_map(sheet: Worksheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in sheet[1]:
        if isinstance(cell.value, str) and isinstance(cell.column, int):
            headers[cell.value.strip()] = cell.column
    if "用例ID" not in headers:
        raise RuntimeError(f"{sheet.title}缺少用例ID字段")
    return headers


def _find_case_row(sheet: Worksheet, headers: dict[str, int], case_id: str) -> int:
    for row in range(2, sheet.max_row + 1):
        current = str(sheet.cell(row, headers["用例ID"]).value or "").strip()
        if current == case_id:
            return row
    raise RuntimeError(f"Excel中找不到用例ID：{case_id}")


def _ensure_history_sheet(workbook) -> Worksheet:
    if HISTORY_SHEET in workbook.sheetnames:
        return workbook[HISTORY_SHEET]
    sheet = workbook.create_sheet(HISTORY_SHEET)
    sheet.append(HISTORY_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = (22, 20, 22, 22, 14, 16, 60, 60)
    for cell, width in zip(sheet[1], widths, strict=True):
        sheet.column_dimensions[cell.column_letter].width = width
    sheet.freeze_panes = "A2"
    return sheet

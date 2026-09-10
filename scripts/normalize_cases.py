from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from yanjia_automation.excel.android_catalog import (
    ANDROID_STEP_SHEET_CASE_IDS,
    SINGLE_SHEET_SOURCE,
    get_android_case_definition,
)

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "test_case.xlsx"
PACKAGE = "com.xiaofutech.yanjia_ai"

# node, locator, action, verification, assertion
AUTOMATED = {
    "TC-LOGIN-001": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "login_username_et,login_pwd_et,login_tv,cover_prompt_cl",
        "input,input,click,verify",
        "错误账号登录后显示失败提示",
        "element_visible",
    ),
    "TC-LOGIN-002": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "login_username_et,login_pwd_et,login_tv,cover_prompt_desc_tv",
        "input,input,click,verify",
        "错误密码登录后显示账号密码错误提示",
        "text_contains",
    ),
    "TC-LOGIN-003": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "login_username_et,login_pwd_et,login_tv",
        "clear,clear,click,verify",
        "空账号密码不会离开 LoginActivity",
        "activity_endswith",
    ),
    "TC-LOGIN-007": (
        "tests/smoke/test_read_only_smoke.py::test_login_and_home_contract",
        "login_username_et,login_pwd_et,login_tv,login_store_rv,main_fbl",
        "input,input,click,click,verify",
        "MainActivity 且首页 main_fbl 可见",
        "activity_endswith+element_visible",
    ),
    "TC-HOME-001": (
        "tests/smoke/test_read_only_smoke.py::test_login_and_home_contract",
        "main_records_ll,main_meiji_ll,main_case_ll,main_set_cl",
        "verify",
        "首页四个入口均可见",
        "elements_visible",
    ),
    "TC-HOME-008": (
        "tests/smoke/test_read_only_smoke.py::test_customer_list_contract",
        "main_records_ll,customer_records_rv",
        "click,verify",
        "CustomerRecordsActivity 且列表可见",
        "activity_endswith+element_visible",
    ),
    "TC-HOME-009": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "main_search_tv",
        "click,verify",
        "搜索入口进入 CustomerRecordsActivity",
        "activity_endswith",
    ),
    "TC-HOME-006": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "main_search_et,main_search_tv,a_records_cl,a_records_unique_tv",
        "input,click,verify,verify",
        "进入顾客列表且至少一个结果手机号包含搜索片段",
        "activity_endswith+element_count_gte+text_contains",
    ),
    "TC-HOME-007": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "main_search_et,main_search_tv,a_records_cl,a_records_unique_tv",
        "input,click,verify,verify",
        "进入顾客列表，仅一个结果且手机号精确匹配受控配置变量",
        "activity_endswith+element_count_equals+text_equals",
    ),
    "TC-CUSTOMER-005": (
        "tests/smoke/test_read_only_smoke.py::test_customer_list_contract",
        "a_records_cl,a_records_name_tv,a_records_age_tv,a_records_unique_tv,a_records_time_tv,a_records_count_tv",
        "verify",
        "至少一张顾客卡片且关键字段控件存在",
        "element_count+elements_visible",
    ),
    "TC-CUSTOMER-006": (
        "tests/smoke/test_read_only_smoke.py::test_first_customer_detail_contract",
        "a_records_cl,customer_detail_name_tv,customer_detail_info_tv,customer_detail_edit_tv",
        "click,verify",
        "CustomerDetailActivity 且详情唯一元素可见",
        "activity_endswith+elements_visible",
    ),
    "TC-CUSTOMER-007": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "main_records_ll,customer_records_rv,customer_records_date_ll,customer_records_date_cl,date_picker_year_wheel,date_picker_month_wheel,date_picker_day_wheel,customer_records_date_confirm_tv,customer_records_date_tv,empty_tv",
        "click,verify,click,verify,select_date_range,click,verify",
        "日期范围显示为2015-02-04~2015-02-13且筛选结果为空",
        "text_equals+element_visible",
    ),
    "TC-HOME-010": (
        "tests/smoke/test_read_only_smoke.py::test_case_library_contract",
        "main_case_ll,case_rv",
        "click,verify",
        "CaseActivity 且案例列表可见",
        "activity_endswith+element_visible",
    ),
    "TC-CASE-001": (
        "tests/smoke/test_read_only_smoke.py::test_case_library_contract",
        "case_search_et,a_case_tag_tv,a_case_category_tv,case_rv",
        "verify",
        "案例搜索、标签、类别和列表控件存在",
        "elements_visible",
    ),
    "TC-CASE-002": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "main_case_ll,case_search_et,case_search_tv,empty_tv,a_case_image_cl",
        "click,input,click,verify",
        "不存在的标签搜索后显示空结果且没有案例卡片",
        "activity_endswith+element_visible+element_not_visible",
    ),
    "TC-CASE-003": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "main_case_ll,case_search_et,case_search_tv,a_case_image_cl",
        "click,input,click,verify",
        "标签搜索后至少一张案例卡片包含指定标签",
        "activity_endswith+element_count",
    ),
    "TC-CASE-004": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "main_case_ll,case_tag_expand_tv,a_case_tag_tv,a_case_image_cl",
        "click,click,click,verify",
        "展开并单选标签后至少一张案例卡片包含指定标签",
        "activity_endswith+element_count",
    ),
    "TC-HOME-011": (
        "tests/smoke/test_read_only_smoke.py::test_profile_contract",
        "main_set_cl,set_logout_tv",
        "click,verify",
        "SetActivity 且设置页根控件可见",
        "activity_endswith+element_visible",
    ),
    "TC-HOME-012": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "main_meiji_ll",
        "click,verify",
        "美际学院入口进入 WebViewPCActivity",
        "activity_endswith",
    ),
    "TC-CUSTOMER-020": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "a_records_cl,a_records_head_ifv,a_records_name_tv,a_records_age_tv,a_records_unique_tv,a_records_time_tv,a_records_count_tv",
        "verify",
        "首张顾客卡片关键字段控件均存在",
        "element_count+elements_visible",
    ),
    "TC-CUSTOMER-021": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "a_records_cl,customer_detail_name_tv",
        "click,verify",
        "进入 CustomerDetailActivity 且详情页可见",
        "activity_endswith+element_visible",
    ),
    "TC-CUSTOMER-022": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_other_tv,customer_records_other_cl,customer_records_other_sex_man_tv,customer_records_other_age_1_tv,customer_records_other_remark_et,customer_records_other_reset_tv",
        "click,verify,click,verify,click,verify,input,verify,click,verify",
        "重置后性别和年龄取消选中，备注恢复搜索备注记录，筛选面板仍打开",
        "element_visible+element_selected+attribute_equals",
    ),
    "TC-DETAIL-001": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_username_et,customer_edit_save_tv",
        "input,click,verify,click,clear,click,verify",
        "专用顾客唯一匹配，姓名为空时显示姓名不能为空且恢复原资料",
        "element_count_equals+text_contains+restore_verified",
    ),
    "TC-DETAIL-003": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_username_et,customer_edit_save_tv",
        "input,click,verify,click,input,click,verify",
        "专用顾客唯一匹配，33字符姓名显示客户名称长度不合法且恢复原资料",
        "element_count_equals+text_contains+restore_verified",
    ),
    "TC-DETAIL-004": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_sex_tv,customer_edit_sex_man_tv,customer_edit_save_tv",
        "input,click,verify,click,click,click,click,click,verify",
        "专用顾客唯一匹配，保存性别男后重开编辑页校验并恢复原资料",
        "element_count_equals+text_equals+restore_verified",
    ),
    "TC-DETAIL-005": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_sex_tv,customer_edit_sex_female_tv,customer_edit_save_tv",
        "input,click,verify,click,click,click,click,click,verify",
        "专用顾客唯一匹配，保存性别女后重开编辑页校验并恢复原资料",
        "element_count_equals+text_equals+restore_verified",
    ),
    "TC-DETAIL-007": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_birthday_tv,customer_edit_save_tv,date_picker_year_wheel,date_picker_month_wheel,date_picker_day_wheel",
        "input,click,verify,click,select_birthday,click,click,verify",
        "专用顾客唯一匹配，保存生日2022-08-04后重开校验并恢复八项原资料",
        "element_count_equals+text_equals+restore_verified",
    ),
    "TC-DETAIL-008": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_address_et,customer_edit_save_tv",
        "input,click,verify,click,clear,click,click,verify",
        "专用顾客唯一匹配，清空地址保存后重开校验并恢复八项原资料",
        "element_count_equals+attribute_equals+restore_verified",
    ),
    "TC-DETAIL-009": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_address_et,customer_edit_save_tv",
        "input,click,verify,click,input,click,click,verify",
        "专用顾客唯一匹配，特殊字符地址保存后重开精确校验并恢复八项原资料",
        "element_count_equals+text_equals+restore_verified",
    ),
    "TC-DETAIL-010": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_address_et,customer_edit_save_tv",
        "input,click,verify,click,input,click,click,verify",
        "专用顾客唯一匹配，单空格地址保存后重开校验为空并恢复八项原资料",
        "element_count_equals+attribute_equals+restore_verified",
    ),
    "TC-DETAIL-011": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_unique_et,customer_edit_save_tv",
        "input,click,verify,click,clear,click,verify",
        "专用顾客唯一匹配，手机号为空时显示手机号不能为空且恢复原资料",
        "element_count_equals+text_contains+restore_verified",
    ),
    "TC-DETAIL-013": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_remark_et,customer_edit_save_tv",
        "input,click,verify,click,scroll_profile_remark,clear,click,click,scroll_profile_remark,verify",
        "专用顾客唯一匹配，清空个人备注保存后重开校验并恢复八项原资料",
        "element_count_equals+attribute_equals+restore_verified",
    ),
    "TC-DETAIL-014": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_remark_et,customer_edit_save_tv",
        "input,click,verify,click,scroll_profile_remark,input,click,click,scroll_profile_remark,verify",
        "专用顾客唯一匹配，特殊字符个人备注保存后重开精确校验并恢复八项原资料",
        "element_count_equals+text_equals+restore_verified",
    ),
    "TC-DETAIL-015": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_edit_tv,customer_edit_marital_tv,customer_edit_marital_secret_tv,customer_edit_save_tv",
        "input,click,verify,click,click,click,click,click,verify",
        "专用顾客唯一匹配，保存婚姻状态保密后重开校验并恢复八项原资料",
        "element_count_equals+text_equals+restore_verified",
    ),
    "TC-DETAIL-016": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_records_detail_all_remark_ifv,a_records_remark_tag_add_ll,records_remark_tag_cl,records_remark_tag_et,records_remark_tag_save_tv,a_records_remark_tag_tv,a_records_remark_tag_delete_ifv,cover_prompt_right_tv",
        "input,click,verify,click,click,input,click,verify,click,click,verify",
        "专用顾客唯一匹配，新增本轮唯一影像标签并显示；退出受控会话时删除差集标签并恢复原集合",
        "element_count_equals+text_contains+restore_verified",
    ),
    "TC-DETAIL-017": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_records_detail_all_remark_ifv,a_records_remark_tag_add_ll,records_remark_tag_cl,records_remark_tag_save_tv,records_remark_tag_cancel_tv,records_remark_close_ifv",
        "input,click,verify,click,click,click,verify,click,click,verify",
        "空标签点击保存显示请输入标签内容，随后取消弹窗且不创建标签",
        "element_count_equals+page_source_contains+activity_endswith",
    ),
    "TC-DETAIL-018": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_records_detail_all_remark_ifv,a_records_remark_tag_add_ll,records_remark_tag_cl,records_remark_tag_et,records_remark_tag_save_tv,a_records_remark_tag_tv,a_records_remark_tag_delete_ifv,cover_prompt_right_tv",
        "input,click,verify,click,click,input,click,verify,click,click,verify",
        "专用顾客唯一匹配，新增带特殊字符的本轮唯一影像标签并显示；退出时精确恢复原集合",
        "element_count_equals+page_source_contains+restore_verified",
    ),
    "TC-DETAIL-019": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_records_detail_all_remark_ifv,a_records_remark_tag_add_ll,records_remark_tag_cl,records_remark_tag_et,records_remark_tag_save_tv,records_remark_tag_cancel_tv,records_remark_close_ifv",
        "input,click,verify,click,click,input,click,verify,click,click,verify",
        "影像标签输入单个空格后保存显示请输入标签内容，随后取消弹窗且标签集合不变",
        "element_count_equals+page_source_contains+activity_endswith",
    ),
    "TC-DETAIL-023": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "a_records_detail_all_remark_ifv,records_remark_remark_et,records_remark_remark_commit_tv,a_records_remark_list_content_tv",
        "click,input,verify",
        "影像备注输入单个空格后点击提交，备注记录数量保持不变",
        "element_count_unchanged_after_click",
    ),
    "TC-DETAIL-024": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_detail_manager_tv,a_records_detail_all_select_ifv,customer_detail_delete_tv",
        "click,verify,click,verify",
        "管理/完成按钮切换且影像选择和删除控件仅在管理态出现",
        "text_equals+element_visible+element_not_visible",
    ),
    "TC-DETAIL-029": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,customer_detail_manager_tv,a_records_detail_all_select_ifv,customer_detail_delete_tv",
        "input,click,verify,click,verify,click,verify,click,verify",
        "管理态不勾选影像直接点击删除，页面中心显示请选择需要删除的影像",
        "element_count_equals+page_source_contains+element_not_visible",
    ),
    "TC-DETAIL-030": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_records_detail_all_head_ifv,a_records_detail_all_time_tv,a_consultation_result_time_2_tv,skin_result_vp2,f_skin_result_info_view_case_tv,skin_result_back_ll,cover_prompt_v2_tv",
        "input,click,verify,click,open_first_unlinked_image,verify,click,verify,click,verify_selected_image_consultation",
        "唯一种子顾客中动态选择首张未关联影像，保存并退出后所选检测时间出现在咨询单卡片中；新增咨询单保留",
        "element_count_equals+activity_endswith+text_equals+selected_image_consultation",
    ),
    "TC-DETAIL-031": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_consultation_result_remark_ifv,a_records_remark_tag_add_ll,records_remark_tag_cl,records_remark_tag_et,records_remark_tag_save_tv,a_records_remark_tag_tv,a_records_remark_tag_delete_ifv,cover_prompt_right_tv",
        "input,click,verify,click,click,input,click,verify,click,click,verify",
        "专用顾客唯一匹配，新增本轮唯一咨询单标签并显示；退出受控会话时删除差集标签并恢复原集合",
        "element_count_equals+page_source_contains+restore_verified",
    ),
    "TC-DETAIL-032": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_consultation_result_cl,a_consultation_result_remark_ifv,a_records_remark_tag_add_ll,records_remark_tag_cl,records_remark_tag_et,records_remark_tag_save_tv,records_remark_tag_cancel_tv,records_remark_close_ifv",
        "input,click,verify,click,click,input,click,verify,click,click,verify",
        "咨询单标签输入单个空格后保存显示请输入标签内容，随后取消弹窗且标签集合不变",
        "element_count_gte+page_source_contains+activity_endswith",
    ),
    "TC-DETAIL-033": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_consultation_result_remark_ifv,a_records_remark_tag_add_ll,records_remark_tag_cl,records_remark_tag_et,records_remark_tag_save_tv,a_records_remark_tag_tv,a_records_remark_tag_delete_ifv,cover_prompt_right_tv",
        "input,click,verify,click,click,input,click,verify,click,click,verify",
        "专用顾客唯一匹配，新增带特殊字符的本轮唯一咨询单标签并显示；退出时精确恢复原集合",
        "element_count_equals+page_source_contains+restore_verified",
    ),
    "TC-DETAIL-034": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,customer_records_search_tv,a_records_cl,a_consultation_result_cl,a_consultation_result_remark_ifv,a_records_remark_tag_add_ll,records_remark_tag_cl,records_remark_tag_save_tv,records_remark_tag_cancel_tv,records_remark_close_ifv",
        "input,click,verify,click,click,click,verify,click,click,verify",
        "咨询单空标签保存显示请输入标签内容，随后取消弹窗且标签集合不变",
        "element_count_gte+page_source_contains+activity_endswith",
    ),
    "TC-DETAIL-036": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "a_consultation_result_remark_ifv,records_remark_remark_et,records_remark_remark_commit_tv,a_records_remark_list_content_tv",
        "click,input,verify",
        "咨询单备注输入单个空格后点击提交，备注记录数量保持不变",
        "element_count_unchanged_after_click",
    ),
    "TC-PROFILE-001": (
        "tests/smoke/test_read_only_smoke.py::test_profile_contract",
        "set_personal_ll,personal_account_tv",
        "click,verify",
        "账号字段非空；实际值不写入报告",
        "text_not_empty",
    ),
    "TC-PROFILE-002": (
        "tests/smoke/test_read_only_smoke.py::test_profile_contract",
        "personal_name_tv",
        "verify",
        "姓名字段非空；实际值不写入报告",
        "text_not_empty",
    ),
    "TC-PROFILE-003": (
        "tests/smoke/test_read_only_smoke.py::test_profile_contract",
        "personal_register_time_tv",
        "verify",
        "注册时间字段非空；实际值不写入报告",
        "text_not_empty",
    ),
    "TC-PROFILE-004": (
        "tests/smoke/test_read_only_smoke.py::test_profile_contract",
        "personal_phone_tv",
        "verify",
        "手机号字段存在，值按测试账号规则校验",
        "element_visible",
    ),
    "TC-IMAGE-001": (
        "tests/regression/test_seeded_image.py::test_existing_image_opens_viewer",
        "a_records_detail_all_head_ifv,skin_result_back_ll,skin_result_vp2",
        "click,verify",
        "SkinResultActivity 且影像 ViewPager 可见",
        "activity_endswith+element_visible",
    ),
    "TC-IMAGE-030": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "f_skin_result_info_comprehensive_rv,a_skin_result_info_radar_child_score_tv,a_skin_result_info_feature_child_score_tv",
        "verify",
        "AI综合分析评分区可见，炎敏/肤色/色素/亮度/红区五项评分均非空且不记录具体值",
        "element_visible+text_not_empty",
    ),
    "TC-IMAGE-032": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "customer_records_search_et,a_records_detail_all_head_ifv,f_skin_result_info_view_case_tv,case_rv",
        "input,click,click,verify",
        "唯一受控顾客的影像可跳转 CaseActivity，且案例列表可见",
        "activity_endswith+element_visible",
    ),
    "TC-IMAGE-033": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "f_skin_result_info_menu_ifv,f_skin_result_info_menu_report_ll,skin_report_pdfv,skin_report_share_tv",
        "click,click,verify",
        "影像报告跳转 SkinReportActivity，报告预览和分享入口可见",
        "activity_endswith+elements_visible",
    ),
    "TC-IMAGE-034": (
        "tests/excel/test_excel_cases.py::test_excel_case",
        "skin_result_back_ll,cover_prompt_v1_tv,customer_detail_name_tv",
        "click,verify,click,verify",
        "选择不保存直接退出后返回 CustomerDetailActivity",
        "text_equals+activity_endswith+element_visible",
    ),
}

INPUTS = {
    "TC-LOGIN-001": "${INVALID_USERNAME}|${YANJIA_PASSWORD}",
    "TC-LOGIN-002": "${YANJIA_USERNAME}|${INVALID_PASSWORD}",
    "TC-LOGIN-004": "${YANJIA_USERNAME}|<EMPTY>",
    "TC-LOGIN-005": "${YANJIA_USERNAME}|${YANJIA_PASSWORD}|<NO_STORE>",
    "TC-LOGIN-007": "${YANJIA_USERNAME}|${YANJIA_PASSWORD}|${YANJIA_STORE_OR_FIRST}",
    "TC-HOME-002": "${NON_EXISTENT_CUSTOMER_QUERY}",
    "TC-HOME-003": "${NON_EXISTENT_PHONE_QUERY}",
    "TC-HOME-004": "${SEEDED_CUSTOMER_NAME_PART}",
    "TC-HOME-005": "${SEEDED_CUSTOMER_NAME}",
    "TC-HOME-006": "186",
    "TC-HOME-007": "${SEEDED_CUSTOMER_PHONE}",
    "TC-CUSTOMER-001": "${EMPTY_DATE_RANGE}",
    "TC-CUSTOMER-002": "${SEEDED_DATE_RANGE}",
    "TC-CUSTOMER-004": "${SEEDED_CUSTOMER_UNDER_18}",
    "TC-CUSTOMER-007": "2015-02-04~2015-02-13",
    "TC-CUSTOMER-022": "123",
    "TC-IMAGE-001": "${SEEDED_CUSTOMER_WITH_IMAGE}",
    "TC-IMAGE-002": "${SEEDED_CUSTOMER_WITH_IMAGE}",
    "TC-IMAGE-003": "${SEEDED_CUSTOMER_WITH_IMAGE}",
    "TC-IMAGE-004": "${SEEDED_CUSTOMER_WITH_HISTORY_IMAGES}",
    "TC-IMAGE-005": "${SEEDED_CUSTOMER_WITH_HISTORY_IMAGES}",
    "TC-IMAGE-030": "咨询单特定",
    "TC-IMAGE-032": "咨询单特定",
    "TC-IMAGE-033": "咨询单特定",
    "TC-IMAGE-034": "咨询单特定",
    "TC-CASE-002": "${NON_EXISTENT_CASE_TAG}",
    "TC-CASE-003": "火",
    "TC-CASE-004": "火",
    "TC-DETAIL-007": "2022-08-04",
    "TC-DETAIL-009": "……&&*……*&",
    "TC-DETAIL-010": "${SPACE}",
    "TC-DETAIL-012": "${YANJIA_MUTATION_CUSTOMER_QUERY}|测试备注${RUN_TOKEN}",
    "TC-DETAIL-014": "……&&*……*&",
    "TC-DETAIL-015": "保密",
    "TC-DETAIL-016": "${YANJIA_MUTATION_CUSTOMER_QUERY}|自动化标签${RUN_TOKEN}",
    "TC-DETAIL-018": "${YANJIA_MUTATION_CUSTOMER_QUERY}|%……&*${RUN_TOKEN}",
    "TC-DETAIL-019": "${SPACE}",
    "TC-DETAIL-020": "${YANJIA_MUTATION_CUSTOMER_QUERY}|自动化删除标签${RUN_TOKEN}",
    "TC-DETAIL-021": "${YANJIA_MUTATION_CUSTOMER_QUERY}|影像备注${RUN_TOKEN}",
    "TC-DETAIL-022": "${YANJIA_MUTATION_CUSTOMER_QUERY}|……&&*……*&${RUN_TOKEN}",
    "TC-DETAIL-026": "${YANJIA_MUTATION_CUSTOMER_QUERY}|自动化${RUN_TOKEN}",
    "TC-DETAIL-027": "${YANJIA_MUTATION_CUSTOMER_QUERY}|${RUN_PHONE}",
    "TC-DETAIL-030": "咨询单特定",
    "TC-DETAIL-031": "${YANJIA_MUTATION_CUSTOMER_QUERY}|咨询单标签${RUN_TOKEN}",
    "TC-DETAIL-032": "咨询单特定|${SPACE}",
    "TC-DETAIL-033": "${YANJIA_MUTATION_CUSTOMER_QUERY}|%……&*${RUN_TOKEN}",
    "TC-DETAIL-034": "咨询单特定|",
    "TC-DETAIL-035": "${YANJIA_MUTATION_CUSTOMER_QUERY}|咨询单备注${RUN_TOKEN}",
    "TC-DETAIL-037": "${YANJIA_MUTATION_CUSTOMER_QUERY}|自动化删除备注${RUN_TOKEN}",
}

DESTRUCTIVE = {
    "TC-DETAIL-020",
    "TC-DETAIL-025",
    "TC-DETAIL-028",
    "TC-DETAIL-037",
    "TC-CASE-005",
}
MUTATING = {
    "TC-DETAIL-001",
    "TC-DETAIL-002",
    "TC-DETAIL-003",
    "TC-DETAIL-004",
    "TC-DETAIL-005",
    "TC-DETAIL-006",
    "TC-DETAIL-007",
    "TC-DETAIL-008",
    "TC-DETAIL-009",
    "TC-DETAIL-010",
    "TC-DETAIL-011",
    "TC-DETAIL-012",
    "TC-DETAIL-013",
    "TC-DETAIL-014",
    "TC-DETAIL-015",
    "TC-DETAIL-016",
    "TC-DETAIL-018",
    "TC-DETAIL-021",
    "TC-DETAIL-022",
    "TC-DETAIL-026",
    "TC-DETAIL-027",
    "TC-DETAIL-030",
    "TC-DETAIL-031",
    "TC-DETAIL-033",
    "TC-DETAIL-035",
}
AUTH_NEGATIVE = {f"TC-LOGIN-{number:03d}" for number in range(1, 6)}
MODULES = {
    "账号登录": "login",
    "首页": "home",
    "首页搜索": "home",
    "首页跳转": "home",
    "顾客列表": "customer",
    "顾客详情": "detail",
    "影像阅览": "image",
    "案例库": "case",
    "个人中心": "profile",
}


def main() -> None:
    _validate_automation_catalog()
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["自动化测试用例"]
    headers: dict[str, int] = {}
    for cell in sheet[1]:
        if isinstance(cell.value, str) and isinstance(cell.column, int):
            headers[cell.value] = cell.column
    for name in ("自动化状态", "pytest节点", "标签", "移动端备注"):
        if name not in headers:
            headers[name] = sheet.max_column + 1
            target = sheet.cell(1, headers[name], name)
            _copy_style(sheet.cell(1, headers["实际结果"]), target)
            sheet.column_dimensions[target.column_letter].width = 28

    row_by_id: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        case_id = str(_get(sheet, row, headers, "用例ID") or "").strip()
        if not case_id:
            continue
        row_by_id[case_id] = row
        module = str(_get(sheet, row, headers, "模块") or "").strip()
        priority = str(_get(sheet, row, headers, "优先级") or "").upper()
        _set(sheet, row, headers, "优先级", priority)
        _set(sheet, row, headers, "实际结果", "NOT_RUN")
        _normalize_timeout(sheet, row, headers)

        if case_id in INPUTS:
            _set(sheet, row, headers, "输入数据", INPUTS[case_id])
        if case_id.startswith("TC-PROFILE-") or case_id == "TC-CASE-001":
            _set(sheet, row, headers, "输入数据", "")
        if case_id in {"TC-HOME-006", "TC-HOME-007", "TC-CUSTOMER-022"}:
            _set(sheet, row, headers, "数据类型", "string")
        if case_id in {"TC-LOGIN-001", "TC-LOGIN-002", "TC-LOGIN-007"}:
            _set(sheet, row, headers, "数据类型", "string|string")

        status, node, note = "BACKLOG", "", "待按真实页面补充 Appium 定位与断言"
        definition = get_android_case_definition(case_id)
        if case_id in AUTOMATED:
            node, locators, actions, verification, assertion = AUTOMATED[case_id]
            status = "AUTOMATED_EXTENDED" if case_id == "TC-IMAGE-001" else "AUTOMATED"
            _set(sheet, row, headers, "元素定位器", _ids(locators))
            _set(sheet, row, headers, "操作类型", actions)
            _set(sheet, row, headers, "验证点", verification)
            _set(sheet, row, headers, "断言类型", assertion)
            note = "已根据 Android 16 平板真实 resource-id 实现"
        elif definition is not None and definition.supports(SINGLE_SHEET_SOURCE):
            status = definition.automation_status_for(SINGLE_SHEET_SOURCE)
            note = "已登记在 Android 单表目录；运行时由单表加载器生成 Appium 步骤"
        elif case_id in DESTRUCTIVE:
            status, note = "DISABLED_DESTRUCTIVE", "只能删除本轮创建且可精确恢复的数据"
        elif case_id == "TC-DETAIL-002":
            status, note = "NEEDS_PRODUCT_RULE", "特殊字符应保存还是拦截需要产品确认"
        elif case_id == "TC-DETAIL-006":
            status = "BLOCKED_PRODUCT"
            note = "真实 Android 顾客编辑页只有男/女选项，没有保密性别选项"

        _set(sheet, row, headers, "自动化状态", status)
        _set(sheet, row, headers, "pytest节点", node)
        _set(sheet, row, headers, "标签", ",".join(_tags(case_id, module, priority)))
        _set(sheet, row, headers, "移动端备注", note)

    _fix_cases(sheet, headers, row_by_id)
    _update_references(workbook)
    _write_notes(workbook)
    workbook.save(WORKBOOK)
    print("Normalized test_case.xlsx without reading env.txt or copying credential values.")


def _validate_automation_catalog() -> None:
    """Fail early if the canonical step metadata drifts from the catalog."""

    actual = frozenset(AUTOMATED)
    extra = actual - ANDROID_STEP_SHEET_CASE_IDS
    missing_non_migrated = {
        case_id
        for case_id in ANDROID_STEP_SHEET_CASE_IDS - actual
        if not (
            (definition := get_android_case_definition(case_id)) is not None
            and definition.supports(SINGLE_SHEET_SOURCE)
        )
    }
    if extra or missing_non_migrated:
        raise RuntimeError(
            "normalize_cases.py 与 Android 用例目录不一致；"
            f"非单表迁移目录缺少={sorted(missing_non_migrated)}，"
            f"脚本多出={sorted(extra)}"
        )


def _fix_cases(sheet, headers: dict[str, int], rows: dict[str, int]) -> None:
    reset_row = rows["TC-CUSTOMER-022"]
    _set(
        sheet,
        reset_row,
        headers,
        "前置条件",
        "已登录并进入顾客列表页；其他筛选面板初始为未选择，备注框显示占位文本",
    )
    _set(
        sheet,
        reset_row,
        headers,
        "操作步骤",
        "1.打开其他筛选 2.选择男和18-25岁并输入备注123 3.逐项校验已设置 4.点击重置",
    )
    _set(sheet, reset_row, headers, "输入数据", "123")
    _set(
        sheet,
        reset_row,
        headers,
        "期望结果",
        "性别和年龄取消选中，备注恢复搜索备注记录，其他筛选面板仍保持打开",
    )
    _set(
        sheet,
        reset_row,
        headers,
        "验证点",
        "男和18-25岁 selected=false；备注 text=搜索备注记录；筛选面板可见",
    )

    consultation_row = rows["TC-DETAIL-030"]
    _set(
        sheet,
        consultation_row,
        headers,
        "前置条件",
        "已登录；“咨询单特定”种子顾客唯一匹配；至少一张历史影像尚未关联咨询单；已授权保留新增咨询单",
    )
    _set(
        sheet,
        consultation_row,
        headers,
        "操作步骤",
        "1.搜索并打开唯一种子顾客 2.按检测时间动态打开首张未关联咨询单的历史影像 "
        "3.返回并选择保存并退出 4.校验所选检测时间出现在咨询单卡片中",
    )
    _set(sheet, consultation_row, headers, "输入数据", "咨询单特定")
    _set(sheet, consultation_row, headers, "数据类型", "string")
    _set(
        sheet,
        consultation_row,
        headers,
        "期望结果",
        "返回顾客详情页，新增咨询单保留，且咨询单卡片显示所选影像的检测时间",
    )
    _set(
        sheet,
        consultation_row,
        headers,
        "验证点",
        "种子顾客唯一匹配；保存并退出；所选检测时间出现在咨询单卡片；不删除新增咨询单",
    )
    _set(
        sheet,
        consultation_row,
        headers,
        "移动端备注",
        "持久化写入用例；需 --run-seeded --allow-mutation "
        "--no-excel-writeback；按用户要求不删除或回滚新增咨询单",
    )

    if "TC-HOME-009" in rows:
        row = rows["TC-HOME-009"]
        _set(sheet, row, headers, "测试场景", "空关键字搜索")
        _set(sheet, row, headers, "测试点", "验证空关键字搜索的产品行为")
        _set(sheet, row, headers, "期望结果", "按产品规则展示默认结果或输入提示")
    image_points = {
        "TC-IMAGE-002": "验证分屏对比功能",
        "TC-IMAGE-003": "验证镜像对比功能",
        "TC-IMAGE-004": "验证历史分屏对比功能",
        "TC-IMAGE-005": "验证历史镜像对比功能",
    }
    for case_id, point in image_points.items():
        _set(sheet, rows[case_id], headers, "测试点", point)
    _set(
        sheet,
        rows["TC-CUSTOMER-006"],
        headers,
        "期望结果",
        "跳转至顾客详情页，详情唯一元素可见",
    )
    _set(sheet, rows["TC-CASE-001"], headers, "期望结果", "进入案例库页面，case_rv 可见")
    _set(
        sheet,
        rows["TC-CASE-002"],
        headers,
        "操作步骤",
        "1.进入案例库 2.输入不存在的案例标签 3.点击搜索 4.校验空结果且无案例卡片",
    )
    _set(
        sheet,
        rows["TC-CASE-002"],
        headers,
        "期望结果",
        "显示空结果提示，且不存在案例卡片",
    )
    _set(sheet, rows["TC-CASE-002"], headers, "超时(秒)", 15)
    for case_id, noun in (("TC-CASE-005", "案例"),):
        row = rows[case_id]
        _set(sheet, row, headers, "测试场景", f"删除本轮创建的指定测试{noun}")
        _set(
            sheet,
            row,
            headers,
            "前置条件",
            f"隔离环境；{noun}由本轮创建且有唯一标识；允许 destructive",
        )
        _set(sheet, row, headers, "期望结果", f"仅指定测试{noun}被删除，其他数据不变")
    birthday_row = rows["TC-DETAIL-007"]
    _set(sheet, birthday_row, headers, "输入数据", "2022-08-04")
    _set(
        sheet,
        birthday_row,
        headers,
        "操作步骤",
        "1.进入顾客资料编辑页 2.选择生日2022-08-04 3.保存 4.重新打开编辑页",
    )
    _set(sheet, birthday_row, headers, "期望结果", "生日显示为2022-08-04")
    _set(sheet, birthday_row, headers, "验证点", "生日字段精确显示2022-08-04")

    address_contracts = {
        "TC-DETAIL-008": ("清空地址并保存", "", "地址字段为空"),
        "TC-DETAIL-009": (
            "输入特殊字符地址并保存",
            "……&&*……*&",
            "地址精确显示……&&*……*&",
        ),
        "TC-DETAIL-010": ("输入单个空格地址并保存", "${SPACE}", "地址字段为空"),
    }
    for case_id, (steps, input_value, expected) in address_contracts.items():
        row = rows[case_id]
        _set(sheet, row, headers, "操作步骤", f"1.进入顾客资料编辑页 2.{steps} 3.重新打开编辑页")
        _set(sheet, row, headers, "输入数据", input_value)
        _set(sheet, row, headers, "期望结果", expected)
        _set(sheet, row, headers, "验证点", expected)

    remark_contracts = {
        "TC-DETAIL-013": ("清空个人备注并保存", "", "个人备注字段为空"),
        "TC-DETAIL-014": (
            "输入特殊字符个人备注并保存",
            "……&&*……*&",
            "个人备注精确显示……&&*……*&",
        ),
    }
    for case_id, (steps, input_value, expected) in remark_contracts.items():
        row = rows[case_id]
        _set(
            sheet,
            row,
            headers,
            "操作步骤",
            f"1.进入顾客资料编辑页 2.滚动到个人备注 3.{steps} 4.重新打开编辑页并滚动到个人备注",
        )
        _set(sheet, row, headers, "输入数据", input_value)
        _set(sheet, row, headers, "期望结果", expected)
        _set(sheet, row, headers, "验证点", expected)

    marital_row = rows["TC-DETAIL-015"]
    _set(
        sheet,
        marital_row,
        headers,
        "操作步骤",
        "1.进入顾客资料编辑页 2.选择婚姻状态保密 3.保存 4.重新打开编辑页",
    )
    _set(sheet, marital_row, headers, "输入数据", "保密")
    _set(sheet, marital_row, headers, "期望结果", "婚姻状态显示为保密")
    _set(sheet, marital_row, headers, "验证点", "婚姻状态字段精确显示保密")

    tag_roundtrip_contracts = {
        "TC-DETAIL-016": (
            "${YANJIA_MUTATION_CUSTOMER_QUERY}|自动化标签${RUN_TOKEN}",
            "专用顾客第一张影像备注",
            "自动化标签${RUN_TOKEN}",
        ),
        "TC-DETAIL-018": (
            "${YANJIA_MUTATION_CUSTOMER_QUERY}|%……&*${RUN_TOKEN}",
            "专用顾客第一张影像备注",
            "%……&*${RUN_TOKEN}",
        ),
        "TC-DETAIL-031": (
            "${YANJIA_MUTATION_CUSTOMER_QUERY}|咨询单标签${RUN_TOKEN}",
            "专用顾客第一条咨询单备注",
            "咨询单标签${RUN_TOKEN}",
        ),
        "TC-DETAIL-033": (
            "${YANJIA_MUTATION_CUSTOMER_QUERY}|%……&*${RUN_TOKEN}",
            "专用顾客第一条咨询单备注",
            "%……&*${RUN_TOKEN}",
        ),
    }
    for case_id, (input_value, entry, label) in tag_roundtrip_contracts.items():
        row = rows[case_id]
        _set(
            sheet,
            row,
            headers,
            "前置条件",
            "已登录；专用顾客唯一匹配；存在目标备注入口；已授权写入并完成本次只读预检",
        )
        _set(
            sheet,
            row,
            headers,
            "操作步骤",
            f"1.进入{entry} 2.记录原标签集合 3.新增标签{label}并保存 "
            "4.校验显示 5.删除本轮差集标签并验证原集合恢复",
        )
        _set(sheet, row, headers, "输入数据", input_value)
        _set(
            sheet,
            row,
            headers,
            "期望结果",
            f"新增标签{label}显示成功，测试结束后原标签集合完全恢复",
        )
        _set(
            sheet,
            row,
            headers,
            "验证点",
            "页面树包含本轮唯一标签；恢复后标签文本、数量和删除控件数量与运行前一致",
        )

    remark_roundtrip_contracts = {
        "TC-DETAIL-012": (
            "${YANJIA_MUTATION_CUSTOMER_QUERY}|测试备注${RUN_TOKEN}",
            "专用顾客第一张影像备注",
            "测试备注${RUN_TOKEN}",
        ),
        "TC-DETAIL-021": (
            "${YANJIA_MUTATION_CUSTOMER_QUERY}|影像备注${RUN_TOKEN}",
            "专用顾客第一张影像备注",
            "影像备注${RUN_TOKEN}",
        ),
        "TC-DETAIL-022": (
            "${YANJIA_MUTATION_CUSTOMER_QUERY}|……&&*……*&${RUN_TOKEN}",
            "专用顾客第一张影像备注",
            "……&&*……*&${RUN_TOKEN}",
        ),
        "TC-DETAIL-035": (
            "${YANJIA_MUTATION_CUSTOMER_QUERY}|咨询单备注${RUN_TOKEN}",
            "专用顾客第一条咨询单备注",
            "咨询单备注${RUN_TOKEN}",
        ),
    }
    for case_id, (input_value, entry, label) in remark_roundtrip_contracts.items():
        row = rows[case_id]
        _set(
            sheet,
            row,
            headers,
            "前置条件",
            "已登录；专用顾客唯一匹配；存在目标备注入口；已授权写入并完成本次只读预检",
        )
        _set(
            sheet,
            row,
            headers,
            "操作步骤",
            f"1.进入{entry} 2.记录原备注集合 3.新增本轮唯一备注{label}并提交 "
            "4.校验显示 5.会话退出时仅删除本轮差集备注并验证原集合恢复",
        )
        _set(sheet, row, headers, "输入数据", input_value)
        _set(
            sheet,
            row,
            headers,
            "期望结果",
            f"新增备注{label}显示成功，测试结束后原备注集合完全恢复",
        )
        _set(
            sheet,
            row,
            headers,
            "验证点",
            "页面树包含本轮唯一备注；恢复后备注文本、记录数和删除控件数与运行前一致",
        )

    card_sync_contracts = {
        "TC-DETAIL-026": ("自动化${RUN_TOKEN}", "姓名"),
        "TC-DETAIL-027": ("${RUN_PHONE}", "手机号"),
    }
    for case_id, (target_value, field_label) in card_sync_contracts.items():
        row = rows[case_id]
        _set(
            sheet,
            row,
            headers,
            "前置条件",
            "已登录；专用顾客唯一匹配；已授权写入与八字段恢复",
        )
        _set(
            sheet,
            row,
            headers,
            "操作步骤",
            f"1.进入专用顾客编辑页 2.修改{field_label}为本轮唯一值 "
            f"3.保存并返回列表 4.校验顾客卡片{field_label}同步 5.恢复八项资料",
        )
        _set(
            sheet,
            row,
            headers,
            "输入数据",
            f"${{YANJIA_MUTATION_CUSTOMER_QUERY}}|{target_value}",
        )
        _set(
            sheet,
            row,
            headers,
            "期望结果",
            f"顾客卡片{field_label}精确显示本轮值；测试结束后八项资料恢复",
        )
        _set(
            sheet,
            row,
            headers,
            "验证点",
            f"卡片{field_label}text_equals；恢复后八项资料逐字段比对",
        )

    empty_tag_row = rows["TC-DETAIL-017"]
    _set(
        sheet,
        empty_tag_row,
        headers,
        "操作步骤",
        "1.进入专用顾客第一张影像备注 2.打开新增标签弹窗 "
        "3.不输入内容点击保存 4.校验提示后取消并返回",
    )
    _set(sheet, empty_tag_row, headers, "输入数据", "")
    _set(sheet, empty_tag_row, headers, "期望结果", "显示请输入标签内容，弹窗保持打开且未创建标签")
    _set(sheet, empty_tag_row, headers, "验证点", "页面树包含请输入标签内容；取消后返回顾客详情")

    tag_validation_contracts = {
        "TC-DETAIL-019": (
            "${SPACE}",
            "进入专用顾客第一张影像备注",
            "输入单个空格标签后点击保存",
        ),
        "TC-DETAIL-032": (
            "咨询单特定|${SPACE}",
            "搜索咨询单种子顾客并进入第一条咨询单备注",
            "输入单个空格标签后点击保存",
        ),
        "TC-DETAIL-034": (
            "咨询单特定|",
            "搜索咨询单种子顾客并进入第一条咨询单备注",
            "不输入标签内容直接点击保存",
        ),
    }
    for case_id, (input_value, entry, validation_action) in tag_validation_contracts.items():
        row = rows[case_id]
        _set(
            sheet,
            row,
            headers,
            "操作步骤",
            f"1.{entry} 2.打开新增标签弹窗 3.{validation_action} 4.校验提示后取消并返回",
        )
        _set(sheet, row, headers, "输入数据", input_value)
        _set(sheet, row, headers, "期望结果", "显示请输入标签内容，弹窗保持打开且标签集合不变")
        _set(sheet, row, headers, "验证点", "页面树包含请输入标签内容；取消后返回顾客详情")


def _update_references(workbook) -> None:
    fields = workbook["字段说明"]
    for row in range(2, fields.max_row + 1):
        if fields.cell(row, 1).value == "元素定位器":
            fields.cell(row, 2, "Appium定位器；优先resource-id，其次accessibility id/UiSelector")
            fields.cell(row, 3, f"id={PACKAGE}:id/main_records_ll")
    assertions = workbook["断言类型说明"]
    assertions.cell(1, 3, "Appium/pytest实现示例")
    for row in range(2, assertions.max_row + 1):
        key = assertions.cell(row, 1).value
        if key in {"url_contains", "url_matches"}:
            assertions.cell(row, 3, "仅WebView；原生页改用Activity与唯一元素")
        elif key == "element_visible":
            assertions.cell(row, 3, "assert element.is_displayed()")
        elif key == "text_not_empty":
            assertions.cell(row, 3, "assert element.text.strip()")
    actions = workbook["操作类型说明"]
    actions.cell(1, 3, "Appium实现示例")
    examples = {
        "click": "driver.find_element(AppiumBy.ID, '...').click()",
        "input": "element.clear(); element.send_keys(value)",
        "verify": "WebDriverWait显式等待后assert",
        "hover": "Android不适用；改为long_click或删除",
        "wait": "WebDriverWait显式等待；禁止固定sleep",
        "nav": "点击页面入口；原生应用不使用page.goto",
        "screenshot": "driver.save_screenshot(path)；注意PII门禁",
    }
    for row in range(2, actions.max_row + 1):
        key = actions.cell(row, 1).value
        if key in examples:
            actions.cell(row, 3, examples[key])


def _write_notes(workbook) -> None:
    notes = (
        workbook["移动端执行说明"]
        if "移动端执行说明" in workbook.sheetnames
        else workbook.create_sheet("移动端执行说明")
    )
    if notes.max_row:
        notes.delete_rows(1, notes.max_row)
    values = [
        ("项目", "内容"),
        ("包名", PACKAGE),
        ("已验证版本", "2.1.9(Build 8) / versionCode 20109"),
        ("执行框架", "Appium 3 + UiAutomator2 + Python + pytest"),
        ("定位优先级", "resource-id > accessibility id > UiSelector > XPath"),
        ("实际结果", "由pytest/JUnit/Allure产生；本表保持NOT_RUN"),
        ("凭据", "只引用环境变量，不保存真实值"),
        ("破坏性", "默认禁用；只能作用于本轮创建且可恢复的数据"),
        ("隐私", "顾客、影像、页面树和截图可能含PII，附件默认关闭"),
    ]
    for row_index, row_values in enumerate(values, 1):
        for column_index, value in enumerate(row_values, 1):
            notes.cell(row_index, column_index, value)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 90
    for cell in notes[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _tags(case_id: str, module: str, priority: str) -> list[str]:
    tags = [priority.lower(), MODULES.get(module, "other"), "tablet"]
    definition = get_android_case_definition(case_id)
    if case_id in DESTRUCTIVE or (
        definition is not None and definition.mutability == "destructive"
    ):
        tags += ["destructive", "serial"]
    elif case_id in MUTATING or (definition is not None and definition.mutability == "mutating"):
        tags += ["mutating", "serial"]
    else:
        tags.append("readonly")
    tags += ["auth_negative", "no_retry"] if case_id in AUTH_NEGATIVE else ["requires_auth"]
    if case_id.startswith(("TC-CUSTOMER", "TC-IMAGE")) or (
        definition is not None and definition.requires_seed
    ):
        tags.append("requires_seed")
    if definition is not None and "restorable" in definition.tags:
        tags.append("restorable")
    if definition is not None and "persistent" in definition.tags:
        tags.append("persistent")
    if case_id in AUTOMATED:
        tags.append("extended" if case_id == "TC-IMAGE-001" else "smoke")
    return list(dict.fromkeys(filter(None, tags)))


def _ids(names: str) -> str:
    return ",".join(f"id={PACKAGE}:id/{name}" for name in names.split(","))


def _normalize_timeout(sheet, row: int, headers: dict[str, int]) -> None:
    value = _get(sheet, row, headers, "超时(秒)")
    if value not in (None, ""):
        try:
            _set(sheet, row, headers, "超时(秒)", int(value))
        except (TypeError, ValueError):
            pass


def _get(sheet, row: int, headers: dict[str, int], name: str):
    return sheet.cell(row, headers[name]).value


def _set(sheet, row: int, headers: dict[str, int], name: str, value) -> None:
    sheet.cell(row, headers[name], value)


def _copy_style(source, target) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format


if __name__ == "__main__":
    main()

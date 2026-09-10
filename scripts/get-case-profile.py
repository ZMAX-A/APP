from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from yanjia_automation.excel.android_catalog import get_android_case_definition
from yanjia_automation.excel.safety import DEDICATED_CUSTOMER_TARGET_CASE_IDS
from yanjia_automation.excel.workbook import CASE_SHEET, TRUE_VALUES

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "test_case.xlsx"


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Return one Excel case's launcher safety profile as JSON."
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--case-id", required=True)
    args = parser.parse_args()

    case_id = args.case_id.strip()
    if not case_id:
        parser.error("--case-id cannot be blank")

    workbook_path = args.workbook.resolve()
    if not workbook_path.is_file():
        parser.error(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if CASE_SHEET not in workbook.sheetnames:
            parser.error(f"Missing worksheet: {CASE_SHEET}")
        sheet = workbook[CASE_SHEET]
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, ())
        headers = {
            str(value or "").strip(): index for index, value in enumerate(header_row)
        }
        required = {"用例ID", "模块", "测试场景", "自动化状态", "标签", "是否执行"}
        missing = sorted(required - set(headers))
        if missing:
            parser.error("Missing columns: " + ", ".join(missing))
        matches = [
            row
            for row in rows
            if str(row[headers["用例ID"]] or "").strip() == case_id
        ]
    finally:
        workbook.close()

    if len(matches) != 1:
        parser.error(f"Expected exactly one case for ID: {case_id}")

    row = matches[0]
    execution_value = str(row[headers["是否执行"]] or "").strip().lower()
    automation_status = str(row[headers["自动化状态"]] or "").strip().upper()
    enabled = (
        execution_value in TRUE_VALUES
        if execution_value
        else automation_status.startswith("AUTOMATED")
    )
    if not enabled:
        parser.error(f"Case is not enabled: {case_id}")

    configured_tags = tuple(
        dict.fromkeys(
            part.strip().lower()
            for part in str(row[headers["标签"]] or "").split(",")
            if part.strip()
        )
    )
    definition = get_android_case_definition(case_id)
    case_tags = configured_tags or (definition.tags if definition is not None else ())
    tags = set(case_tags)
    module = str(row[headers["模块"]] or "").strip()
    scenario = str(row[headers["测试场景"]] or "").strip()
    title = " | ".join(part for part in (case_id, module, scenario) if part)
    profile = {
        "caseId": case_id,
        "title": title,
        "tags": list(case_tags),
        "requiresSeed": "requires_seed" in tags,
        "mutating": "mutating" in tags,
        "destructive": "destructive" in tags,
        "persistent": "persistent" in tags,
        "dedicatedCustomer": case_id in DEDICATED_CUSTOMER_TARGET_CASE_IDS,
    }
    print(json.dumps(profile, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

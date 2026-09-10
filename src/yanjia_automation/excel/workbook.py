from __future__ import annotations

from fnmatch import fnmatchcase
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from yanjia_automation.excel.models import ExcelCase, ExcelStep
from yanjia_automation.excel.single_sheet import (
    ANDROID_SINGLE_SHEET_CASES,
    SINGLE_SHEET_REQUIRED_HEADERS,
    SingleSheetMigrationError,
    load_android_single_sheet_cases,
)

CASE_SHEET = "自动化测试用例"
STEP_SHEET = "自动化执行步骤"

CASE_REQUIRED_HEADERS = {
    "用例ID",
    "模块",
    "测试场景",
    "测试点",
    "优先级",
    "输入数据",
    "期望结果",
    "超时(秒)",
    "自动化状态",
    "标签",
    "是否执行",
}

STEP_REQUIRED_HEADERS = {
    "用例ID",
    "步骤序号",
    "步骤名称",
    "操作类型",
    "元素定位器",
    "输入数据",
    "断言类型",
    "期望值",
    "超时(秒)",
    "元素索引",
    "失败继续",
    "是否执行",
}

TRUE_VALUES = {"1", "true", "yes", "y", "on", "是", "启用", "执行"}
FALSE_VALUES = {"0", "false", "no", "n", "off", "否", "禁用", "不执行"}


class WorkbookFormatError(ValueError):
    """Raised when the Excel workbook cannot be executed safely."""


class ExcelCaseRepository:
    def __init__(self, path: Path) -> None:
        self.path = path.resolve()

    def load_cases(
        self,
        *,
        case_patterns: tuple[str, ...] = (),
        required_tags: tuple[str, ...] = (),
    ) -> list[ExcelCase]:
        if not self.path.is_file():
            raise WorkbookFormatError(f"Excel用例文件不存在：{self.path}")

        # The execution workbook is small, while read-only worksheets make each
        # sheet.cell() lookup rescan worksheet XML. Loading it eagerly keeps the
        # existing validation semantics and reduces targeted collection from
        # minutes to well below one second.
        workbook = load_workbook(self.path, read_only=False, data_only=True)
        try:
            if CASE_SHEET not in workbook.sheetnames:
                raise WorkbookFormatError(f"缺少工作表：{CASE_SHEET}")
            if STEP_SHEET not in workbook.sheetnames:
                case_sheet = workbook[CASE_SHEET]
                _ensure_sheet_dimensions(case_sheet)
                try:
                    case_headers = _headers(case_sheet, SINGLE_SHEET_REQUIRED_HEADERS)
                    cases = load_android_single_sheet_cases(
                        case_sheet,
                        case_headers,
                        case_patterns=case_patterns,
                    )
                except SingleSheetMigrationError as error:
                    raise WorkbookFormatError(str(error)) from error
                selected = [case for case in cases if case.enabled]
                if not selected:
                    raise WorkbookFormatError(
                        "Excel筛选结果为空：没有找到已启用且符合条件的 Android 用例"
                    )
                return selected

            case_sheet = workbook[CASE_SHEET]
            step_sheet = workbook[STEP_SHEET]
            _ensure_sheet_dimensions(case_sheet)
            _ensure_sheet_dimensions(step_sheet)
            case_headers = _headers(case_sheet, CASE_REQUIRED_HEADERS)
            step_headers = _headers(step_sheet, STEP_REQUIRED_HEADERS)
            steps_by_case = self._load_steps(step_sheet, step_headers)
            try:
                single_sheet_headers = _headers(case_sheet, SINGLE_SHEET_REQUIRED_HEADERS)
                migrated_cases = load_android_single_sheet_cases(
                    case_sheet,
                    single_sheet_headers,
                    case_patterns=tuple(sorted(ANDROID_SINGLE_SHEET_CASES)),
                )
            except SingleSheetMigrationError as error:
                raise WorkbookFormatError(str(error)) from error
            migrated_by_id = {case.case_id: case for case in migrated_cases}
            for case_id, migrated_case in migrated_by_id.items():
                steps_by_case.setdefault(case_id, list(migrated_case.steps))
            cases = self._load_case_rows(
                case_sheet,
                case_headers,
                steps_by_case,
                migrated_by_id=migrated_by_id,
            )
        finally:
            workbook.close()

        selected = [case for case in cases if case.enabled]
        if case_patterns:
            selected = [
                case
                for case in selected
                if any(fnmatchcase(case.case_id, pattern) for pattern in case_patterns)
            ]
        if required_tags:
            required = {tag.lower() for tag in required_tags}
            selected = [case for case in selected if required.issubset(set(case.tags))]

        missing_steps = [case.case_id for case in selected if not case.steps]
        if missing_steps:
            joined = ", ".join(missing_steps)
            raise WorkbookFormatError(f"已启用用例缺少执行步骤：{joined}")
        return selected

    def _load_steps(
        self,
        sheet: Worksheet,
        headers: dict[str, int],
    ) -> dict[str, list[ExcelStep]]:
        steps_by_case: dict[str, list[ExcelStep]] = {}
        seen_orders: set[tuple[str, int]] = set()
        for row in range(2, sheet.max_row + 1):
            case_id = _text(_value(sheet, row, headers, "用例ID"))
            if not case_id:
                continue
            order = _integer(_value(sheet, row, headers, "步骤序号"), default=row - 1)
            order_key = case_id, order
            if order_key in seen_orders:
                raise WorkbookFormatError(f"{STEP_SHEET} 第{row}行步骤序号重复：{case_id}/{order}")
            seen_orders.add(order_key)

            action = _text(_value(sheet, row, headers, "操作类型")).lower()
            enabled = _boolean(_value(sheet, row, headers, "是否执行"), default=True)
            if enabled and not action:
                raise WorkbookFormatError(f"{STEP_SHEET} 第{row}行缺少操作类型")

            step = ExcelStep(
                source_row=row,
                case_id=case_id,
                order=order,
                name=_text(_value(sheet, row, headers, "步骤名称")) or f"步骤{order}",
                action=action,
                locator=_optional_text(_value(sheet, row, headers, "元素定位器")),
                input_value=_optional_text(_value(sheet, row, headers, "输入数据")),
                assertion=_optional_lower(_value(sheet, row, headers, "断言类型")),
                expected=_optional_text(_value(sheet, row, headers, "期望值")),
                timeout=_number(_value(sheet, row, headers, "超时(秒)"), default=10.0),
                element_index=_integer(_value(sheet, row, headers, "元素索引"), default=0),
                continue_on_failure=_boolean(
                    _value(sheet, row, headers, "失败继续"), default=False
                ),
                enabled=enabled,
                note=_optional_text(_optional_value(sheet, row, headers, "备注")),
            )
            steps_by_case.setdefault(case_id, []).append(step)

        for steps in steps_by_case.values():
            steps.sort(key=lambda item: (item.order, item.source_row))
        return steps_by_case

    def _load_case_rows(
        self,
        sheet: Worksheet,
        headers: dict[str, int],
        steps_by_case: dict[str, list[ExcelStep]],
        *,
        migrated_by_id: dict[str, ExcelCase],
    ) -> list[ExcelCase]:
        cases: list[ExcelCase] = []
        seen_ids: set[str] = set()
        for row in range(2, sheet.max_row + 1):
            case_id = _text(_value(sheet, row, headers, "用例ID"))
            if not case_id:
                continue
            if case_id in seen_ids:
                raise WorkbookFormatError(f"{CASE_SHEET} 第{row}行用例ID重复：{case_id}")
            seen_ids.add(case_id)

            migrated_case = migrated_by_id.get(case_id)
            automation_status = _text(_value(sheet, row, headers, "自动化状态")) or (
                migrated_case.automation_status if migrated_case is not None else ""
            )
            requested = _boolean(
                _value(sheet, row, headers, "是否执行"),
                default=automation_status.upper().startswith("AUTOMATED"),
            )
            configured_tags = tuple(
                dict.fromkeys(
                    part.strip().lower()
                    for part in _text(_value(sheet, row, headers, "标签")).split(",")
                    if part.strip()
                )
            )
            tags = configured_tags or (
                migrated_case.tags if migrated_case is not None else ()
            )
            enabled_steps = tuple(step for step in steps_by_case.get(case_id, []) if step.enabled)
            cases.append(
                ExcelCase(
                    source_row=row,
                    case_id=case_id,
                    module=_text(_value(sheet, row, headers, "模块")),
                    scenario=_text(_value(sheet, row, headers, "测试场景")),
                    test_point=_text(_value(sheet, row, headers, "测试点")),
                    priority=_text(_value(sheet, row, headers, "优先级")).upper(),
                    precondition=_optional_text(
                        _optional_value(sheet, row, headers, "前置条件")
                    ),
                    input_data=_optional_text(_value(sheet, row, headers, "输入数据")),
                    expected_result=_optional_text(
                        _value(sheet, row, headers, "期望结果")
                    ),
                    timeout=_number(_value(sheet, row, headers, "超时(秒)"), default=10.0),
                    tags=tags,
                    enabled=requested and bool(enabled_steps),
                    automation_status=automation_status,
                    steps=enabled_steps,
                )
            )

        unknown_case_ids = sorted(set(steps_by_case) - seen_ids)
        if unknown_case_ids:
            joined = ", ".join(unknown_case_ids)
            raise WorkbookFormatError(f"执行步骤引用了不存在的用例ID：{joined}")
        return cases


def parse_csv_option(value: str | None) -> tuple[str, ...]:
    if not value:
        return ()
    return tuple(part.strip() for part in value.split(",") if part.strip())


def _ensure_sheet_dimensions(sheet: Worksheet) -> None:
    """Populate bounds for read-only workbooks with no dimension metadata."""
    if sheet.max_row is not None and sheet.max_column is not None:
        return
    calculate_dimension = getattr(sheet, "calculate_dimension", None)
    if calculate_dimension is None:
        return
    try:
        calculate_dimension(force=True)
    except TypeError:
        calculate_dimension()


def _headers(sheet: Worksheet, required: set[str]) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in sheet[1]:
        if isinstance(cell.value, str) and isinstance(cell.column, int):
            headers[cell.value.strip()] = cell.column
    missing = sorted(required - set(headers))
    if missing:
        raise WorkbookFormatError(f"{sheet.title} 缺少字段：{', '.join(missing)}")
    return headers


def _value(sheet: Worksheet, row: int, headers: dict[str, int], name: str) -> Any:
    return sheet.cell(row, headers[name]).value


def _optional_value(
    sheet: Worksheet,
    row: int,
    headers: dict[str, int],
    name: str,
) -> Any:
    column = headers.get(name)
    return sheet.cell(row, column).value if column is not None else None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _optional_text(value: Any) -> str | None:
    text = _text(value)
    return text or None


def _optional_lower(value: Any) -> str | None:
    text = _text(value).lower()
    return text or None


def _boolean(value: Any, *, default: bool) -> bool:
    if value is None or _text(value) == "":
        return default
    normalized = _text(value).lower()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    raise WorkbookFormatError(f"无法识别的是/否值：{value!r}")


def _number(value: Any, *, default: float) -> float:
    if value is None or _text(value) == "":
        return default
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise WorkbookFormatError(f"无法识别的数字：{value!r}") from error
    if result <= 0:
        raise WorkbookFormatError(f"数字必须大于0：{value!r}")
    return result


def _integer(value: Any, *, default: int) -> int:
    if value is None or _text(value) == "":
        return default
    try:
        return int(value)
    except (TypeError, ValueError) as error:
        raise WorkbookFormatError(f"无法识别的整数：{value!r}") from error

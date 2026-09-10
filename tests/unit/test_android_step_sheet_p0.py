from __future__ import annotations

from pathlib import Path
from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from scripts.prepare_excel import CURRENT_STEPS
from yanjia_automation.excel.workbook import ExcelCaseRepository

HEADERS = (
    "用例ID",
    "模块",
    "测试场景",
    "测试点",
    "优先级",
    "前置条件",
    "输入数据",
    "期望结果",
    "超时(秒)",
    "是否执行",
)

STEP_SHEET_PARITY_CASES = (
    ("TC-LOGIN-001", "账号登录", "${INVALID_USERNAME}|${TEST_PASSWORD}"),
    ("TC-LOGIN-002", "账号登录", "${TEST_USERNAME}|${INVALID_PASSWORD}"),
    ("TC-LOGIN-003", "账号登录", None),
    ("TC-LOGIN-004", "账号登录", "${TEST_USERNAME}"),
    ("TC-LOGIN-005", "账号登录", "${TEST_USERNAME}|${TEST_PASSWORD}"),
    ("TC-LOGIN-006", "账号登录", None),
    ("TC-HOME-002", "首页搜索", "不存在的用户xyz123"),
    ("TC-HOME-003", "首页搜索", "100000000000"),
    ("TC-HOME-004", "首页搜索", "t"),
    ("TC-HOME-005", "首页搜索", "余婷"),
    ("TC-HOME-006", "首页搜索", "186"),
    ("TC-HOME-007", "首页搜索", "${SEEDED_CUSTOMER_PHONE}"),
    ("TC-HOME-009", "首页跳转", None),
    ("TC-HOME-012", "首页跳转", None),
    ("TC-CUSTOMER-001", "顾客列表", "不存在的用户xyz123"),
    ("TC-CUSTOMER-002", "顾客列表", "t"),
    ("TC-CUSTOMER-003", "顾客列表", "余婷"),
    ("TC-CUSTOMER-004", "顾客列表", "100000000000"),
    ("TC-CUSTOMER-007", "顾客列表", "2015-02-04~2015-02-13"),
    ("TC-CUSTOMER-008", "顾客列表", "2026-01-01~2026-06-28"),
    ("TC-CUSTOMER-009", "顾客列表", "男"),
    ("TC-CUSTOMER-010", "顾客列表", "女"),
    ("TC-CUSTOMER-012", "顾客列表", "18-25岁"),
    ("TC-CUSTOMER-013", "顾客列表", "26-35岁"),
    ("TC-CUSTOMER-014", "顾客列表", "36-45岁"),
    ("TC-CUSTOMER-015", "顾客列表", "45岁以上"),
    ("TC-CUSTOMER-016", "顾客列表", "18岁以下"),
    ("TC-CUSTOMER-017", "顾客列表", "2015-02-04~2015-02-28"),
    ("TC-CUSTOMER-018", "顾客列表", "2026-01-01~2026-07-30"),
    ("TC-CUSTOMER-020", "顾客列表", None),
    ("TC-CUSTOMER-021", "顾客列表", None),
    ("TC-CUSTOMER-022", "顾客列表", "123"),
    ("TC-DETAIL-001", "顾客详情", None),
    ("TC-DETAIL-011", "顾客详情", None),
    ("TC-DETAIL-023", "顾客详情", None),
    ("TC-DETAIL-024", "顾客详情", None),
    ("TC-DETAIL-029", "顾客详情", "咨询单特定"),
    ("TC-DETAIL-030", "顾客详情", "咨询单特定"),
    ("TC-DETAIL-036", "顾客详情", "咨询单特定| "),
    ("TC-CASE-002", "案例库", "${NON_EXISTENT_CASE_TAG}"),
    ("TC-CASE-003", "案例库", "火"),
    ("TC-CASE-004", "案例库", None),
)

P1_STEP_SHEET_PARITY_CASES = frozenset(
    {
        "TC-LOGIN-004",
        "TC-LOGIN-005",
        "TC-HOME-002",
        "TC-HOME-003",
        "TC-HOME-004",
        "TC-HOME-005",
        "TC-CUSTOMER-001",
        "TC-CUSTOMER-002",
        "TC-CUSTOMER-003",
        "TC-CUSTOMER-004",
        "TC-CUSTOMER-008",
        "TC-CUSTOMER-009",
        "TC-CUSTOMER-010",
        "TC-CUSTOMER-012",
        "TC-CUSTOMER-013",
        "TC-CUSTOMER-014",
        "TC-CUSTOMER-015",
        "TC-CUSTOMER-016",
        "TC-CUSTOMER-017",
        "TC-CUSTOMER-018",
        "TC-CUSTOMER-022",
        "TC-DETAIL-023",
        "TC-DETAIL-024",
        "TC-DETAIL-029",
        "TC-DETAIL-030",
        "TC-DETAIL-036",
        "TC-CASE-003",
        "TC-CASE-004",
    }
)

P2_STEP_SHEET_PARITY_CASES = frozenset({"TC-CUSTOMER-007", "TC-HOME-006"})

P3_STEP_SHEET_PARITY_CASES = frozenset({"TC-HOME-007"})

FIVE_SECOND_PARITY_CASES = frozenset(
    {
        "TC-HOME-002",
        "TC-HOME-003",
        "TC-HOME-004",
        "TC-HOME-005",
        "TC-HOME-006",
        "TC-HOME-007",
    }
)


def test_step_sheet_definitions_match_single_sheet_adapter(tmp_path: Path) -> None:
    workbook_path = tmp_path / "p0-parity.xlsx"
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "自动化测试用例"
    sheet.append(HEADERS)
    for case_id, module, input_data in STEP_SHEET_PARITY_CASES:
        sheet.append(
            (
                case_id,
                module,
                "P0执行源一致性",
                "验证标准步骤表与单表迁移器语义一致",
                (
                    "P3"
                    if case_id in P3_STEP_SHEET_PARITY_CASES
                    else (
                        "P2"
                        if case_id in P2_STEP_SHEET_PARITY_CASES
                        else "P1" if case_id in P1_STEP_SHEET_PARITY_CASES else "P0"
                    )
                ),
                "使用真实 Android 页面契约",
                input_data,
                "生成一致的原生 Appium 步骤",
                5 if case_id in FIVE_SECOND_PARITY_CASES else 10,
                "是",
            )
        )
    workbook.save(workbook_path)
    workbook.close()

    cases = ExcelCaseRepository(workbook_path).load_cases(
        case_patterns=tuple(case_id for case_id, _, _ in STEP_SHEET_PARITY_CASES)
    )
    by_id = {case.case_id: case for case in cases}

    for case_id, _, _ in STEP_SHEET_PARITY_CASES:
        generated = [
            (
                step.action,
                step.locator,
                step.input_value,
                step.assertion,
                step.expected,
                step.timeout,
                step.element_index,
            )
            for step in by_id[case_id].steps
        ]
        canonical = [
            (
                action,
                locator,
                input_value,
                assertion,
                expected,
                float(timeout),
                index,
            )
            for (
                _name,
                action,
                locator,
                input_value,
                assertion,
                expected,
                timeout,
                index,
                _continue_on_failure,
                _enabled,
                _note,
            ) in CURRENT_STEPS[case_id]
        ]
        assert canonical == generated, case_id


def test_canonical_workbook_keeps_single_sheet_migrations_during_transition(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "hybrid.xlsx"
    workbook = Workbook()
    cases = cast(Worksheet, workbook.active)
    cases.title = "自动化测试用例"
    cases.append((*HEADERS, "自动化状态", "标签"))
    cases.append(
        (
            "TC-HOME-009",
            "首页跳转",
            "首页搜索入口",
            "验证标准步骤优先",
            "P0",
            "已登录",
            None,
            "进入顾客列表",
            10,
            "是",
            "AUTOMATED",
            "p0,home,tablet,readonly,requires_auth,smoke",
        )
    )
    cases.append(
        (
            "TC-HOME-002",
            "首页搜索",
            "不存在的姓名",
            "验证单表迁移回退",
            "P1",
            "已登录",
            "不存在的用户xyz123",
            "展示暂无内容",
            10,
            "是",
            "AUTOMATED_ANDROID_MIGRATED",
            "p1,home,tablet,readonly,requires_auth",
        )
    )
    cases.append(
        (
            "TC-CUSTOMER-011",
            "顾客列表",
            "产品规则待确认",
            "没有 Android 步骤",
            "P3",
            "已登录",
            None,
            "待实现",
            10,
            "是",
            "BACKLOG",
            "p3,customer,tablet,readonly,requires_auth",
        )
    )

    steps = workbook.create_sheet("自动化执行步骤")
    steps.append(
        (
            "用例ID",
            "步骤序号",
            "步骤名称",
            "操作类型",
            "元素定位器",
            "输入数据",
            "断言类型",
            "期望值",
            "超时(秒)",
            "元素索引",
            "失败继续",
            "是否执行",
            "备注",
        )
    )
    for order, values in enumerate(CURRENT_STEPS["TC-HOME-009"], 1):
        steps.append(("TC-HOME-009", order, *values))
    workbook.save(workbook_path)
    workbook.close()

    loaded = ExcelCaseRepository(workbook_path).load_cases()
    by_id = {case.case_id: case for case in loaded}

    assert list(by_id) == ["TC-HOME-009", "TC-HOME-002"]
    assert by_id["TC-HOME-009"].steps[1].locator == (
        "id=com.xiaofutech.yanjia_ai:id/main_search_tv"
    )
    assert by_id["TC-HOME-002"].steps[-1].locator == ("id=com.xiaofutech.yanjia_ai:id/empty_tv")
    assert "TC-CUSTOMER-011" not in by_id

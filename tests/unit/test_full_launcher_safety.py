from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tests.unit.test_retry_selection import write_result
from yanjia_automation.excel.workbook import CASE_SHEET

POWERSHELL = shutil.which("powershell")


@pytest.mark.skipif(POWERSHELL is None, reason="Windows PowerShell is required for launcher tests")
@pytest.mark.parametrize(
    ("scenario", "expected_exit", "expected_phase", "expected_remaining"),
    [
        ("readonly_pass", 0, "completed_after_retry", []),
        ("mixed", 1, "completed_with_failures", ["TC-DETAIL-030"]),
        ("mixed_fail", 1, "completed_with_failures", ["TC-DETAIL-030"]),
        ("destructive_pass", 1, "completed_with_failures", ["TC-DETAIL-020"]),
        ("readonly_skip", 1, "completed_with_failures", ["TC-HOME-007"]),
        (
            "restore_failed",
            1,
            "completed_with_failures",
            ["TC-DETAIL-005", "TC-DETAIL-030"],
        ),
    ],
)
def test_full_launcher_aggregates_both_rounds_without_device_or_real_workbook(
    tmp_path: Path,
    scenario: str,
    expected_exit: int,
    expected_phase: str,
    expected_remaining: list[str],
) -> None:
    root = Path(__file__).resolve().parents[2]
    script_dir = tmp_path / "scripts"
    script_dir.mkdir()
    launcher = script_dir / "run-full-visible.ps1"
    source = (root / "scripts" / launcher.name).read_text(encoding="utf-8-sig")
    # Only inject the interpreter location. All control flow and exit handling
    # execute the real launcher; its run-excel dependency is a device-free stub.
    source = source.replace(
        "$python = Join-Path $projectRoot '.venv\\Scripts\\python.exe'",
        "$python = '" + sys.executable.replace("'", "''") + "'",
    )
    launcher.write_text(source, encoding="utf-8-sig")
    shutil.copy2(root / "scripts" / "select_retry_cases.py", script_dir)
    first, retry = tmp_path / "fixtures" / "full", tmp_path / "fixtures" / "retry"
    write_result(first, "TC-HOME-001", "passed")
    write_result(first, "TC-HOME-007", "failed")
    expected_retry_ids = {"TC-HOME-007"}
    expected_blocked_ids: set[str] = set()
    if scenario in {"mixed", "mixed_fail", "restore_failed"}:
        expected_blocked_ids.add("TC-DETAIL-030")
        write_result(
            first, "TC-DETAIL-030", "broken", ("mutating", "persistent"),
            restoration_failed=scenario == "restore_failed",
        )
    if scenario == "restore_failed":
        expected_blocked_ids.add("TC-DETAIL-005")
        write_result(first, "TC-DETAIL-005", "skipped", ("mutating",), mutation_blocked=True)
    if scenario == "destructive_pass":
        expected_blocked_ids.add("TC-DETAIL-020")
        write_result(first, "TC-DETAIL-020", "broken", ("mutating", "destructive"))
    write_result(retry, "TC-HOME-007", "skipped" if scenario == "readonly_skip" else "passed")
    workbook_path = tmp_path / "test_case.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = CASE_SHEET
    sheet.append(["用例ID"])
    all_case_ids = {"TC-HOME-001"} | expected_retry_ids | expected_blocked_ids
    for case_id in sorted(all_case_ids):
        sheet.append([case_id])
    workbook.save(workbook_path)
    workbook.close()
    # Use the actual Excel writer in both stub rounds, including preparation,
    # result styling, history and backups. All files stay in this test directory.
    (script_dir / "write_fixture_results.py").write_text(
        '''import json
import sys
from datetime import datetime
from pathlib import Path
from yanjia_automation.excel.results import ExcelResult, ExcelResultWriter

root = Path(__file__).resolve().parents[1]
run_id = sys.argv[1]
selected = set(sys.argv[2].split(',')) if len(sys.argv) > 2 and sys.argv[2] else None
rows = [json.loads(p.read_text(encoding='utf-8'))
        for p in (root / 'fixtures' / run_id).glob('*-result.json')]
if selected is not None:
    assert {row['name'] for row in rows} == selected
writer = ExcelResultWriter(source_path=root / 'test_case.xlsx', output_path=None,
    run_id=run_id, allure_report_dir=root / 'reports' / 'allure-report' / run_id, enabled=True)
writer.prepare(tuple(row['name'] for row in rows))
for row in rows:
    statuses = {'passed': 'PASS', 'failed': 'FAIL', 'broken': 'ERROR', 'skipped': 'SKIP'}
    status = statuses[row['status']]
    writer.record(ExcelResult(case_id=row['name'], status=status,
        started_at=datetime(2026, 9, 8, 10), finished_at=datetime(2026, 9, 8, 10, 0, 2),
        duration_seconds=2.0 if run_id == 'retry' else 1.0,
        error_message='' if status == 'PASS' else run_id + '-error'))
''',
        encoding="utf-8",
    )
    (script_dir / "run-excel.ps1").write_text(
        r'''param(
    [string]$CaseId,
    [switch]$RunSeeded,
    [switch]$AllowMutation,
    [switch]$AllowDestructive,
    [switch]$CustomerPreflightVerified,
    [switch]$NoWriteBack,
    [switch]$ReadOnlyRetry,
    [switch]$SkipExcelValidation,
    [switch]$VerboseProgress,
    [switch]$KeepAppium,
    [string]$AppiumProcessIdFile
)
$ErrorActionPreference = 'Stop'
$testRoot = Split-Path -Parent $PSScriptRoot
$runName = if ($NoWriteBack) { 'preflight' } elseif ($CaseId) { 'retry' } else { 'full' }
if ($runName -eq 'retry' -and (
    -not $ReadOnlyRetry -or -not $RunSeeded -or $AllowMutation -or
    $AllowDestructive -or $CustomerPreflightVerified
)) {
    throw 'Retry must be limited to read-only cases without mutation authorizations.'
}
if ($runName -ne 'preflight') {
    & $env:YANJIA_UNIT_PYTHON (Join-Path $PSScriptRoot 'write_fixture_results.py') $runName $CaseId
    if ($LASTEXITCODE -ne 0) { throw 'Fixture result writeback failed.' }
}
[System.IO.File]::AppendAllText((Join-Path $testRoot 'calls.txt'), "$runName`n")
$latestDir = Join-Path $testRoot 'reports\allure-report'
New-Item -ItemType Directory -Path $latestDir -Force | Out-Null
if ($runName -ne 'preflight') {
    $resultDir = Join-Path $testRoot "reports\allure-results\$runName"
    New-Item -ItemType Directory -Path $resultDir -Force | Out-Null
    foreach ($file in Get-ChildItem -LiteralPath (Join-Path $testRoot "fixtures\$runName")) {
        Copy-Item -LiteralPath $file.FullName -Destination $resultDir
    }
}
[System.IO.File]::WriteAllText((Join-Path $latestDir 'latest-run.txt'), $runName)
if ($runName -eq 'full') { exit 1 }
exit 0
''',
        encoding="utf-8-sig",
    )
    environment = os.environ.copy()
    environment["PYTHONPATH"] = str(root / "src")
    environment["YANJIA_UNIT_PYTHON"] = sys.executable
    command = "function Start-Sleep {}\n& '" + str(launcher).replace("'", "''") + "'"
    assert POWERSHELL is not None
    completed = subprocess.run(
        [
            POWERSHELL, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
            "-Command", command,
        ],
        cwd=tmp_path, env=environment, capture_output=True, text=True, encoding="utf-8",
        errors="replace", timeout=60,
    )
    assert completed.returncode == expected_exit, completed.stdout + completed.stderr
    status = json.loads(
        (tmp_path / "outputs/launcher-logs/full-run-latest.status.json").read_text(encoding="utf-8")
    )
    assert status["phase"] == expected_phase
    assert status["exitCode"] == expected_exit
    assert sorted(status["remainingFailedCaseIds"]) == sorted(expected_remaining)
    assert set(status["retryCaseIds"]) == expected_retry_ids
    assert set(status["blockedRetryCaseIds"]) == expected_blocked_ids
    assert status["restorationFailed"] is (scenario == "restore_failed")
    calls = (tmp_path / "calls.txt").read_text(encoding="utf-8").splitlines()
    assert calls == ["preflight", "full", "retry"]
    workbook = load_workbook(workbook_path)
    try:
        sheet = workbook[CASE_SHEET]
        headers = {
            cell.value: cell.column for cell in sheet[1] if isinstance(cell.column, int)
        }
        rows = {sheet.cell(row, 1).value: row for row in range(2, sheet.max_row + 1)}
        for case_id in expected_retry_ids:
            row = rows[case_id]
            expected_status = "PASS"
            if scenario == "readonly_skip":
                expected_status = "SKIP"
            assert sheet.cell(row, headers["实际结果"]).value == expected_status
            assert sheet.cell(row, headers["运行编号"]).value == "retry"
            assert sheet.cell(row, headers["执行耗时(秒)"]).value == 2
            assert sheet.cell(row, headers["Allure报告目录"]).value == str(
                tmp_path / "reports/allure-report/retry"
            )
            if expected_status == "PASS":
                assert sheet.cell(row, headers["错误信息"]).value is None
                assert sheet.cell(row, headers["实际结果"]).fill.fgColor.rgb == "FFC6EFCE"
            else:
                assert sheet.cell(row, headers["错误信息"]).value == "retry-error"
        control_row = rows["TC-HOME-001"]
        assert sheet.cell(control_row, headers["运行编号"]).value == "full"
        for case_id in expected_blocked_ids:
            row = rows[case_id]
            assert sheet.cell(row, headers["运行编号"]).value == "full"
        history = list(workbook["执行记录"].values)[1:]
        for case_id in expected_retry_ids:
            assert [row[0] for row in history if row[1] == case_id] == ["full", "retry"]
        for case_id in expected_blocked_ids:
            assert [row[0] for row in history if row[1] == case_id] == ["full"]
    finally:
        workbook.close()

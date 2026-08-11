"""Remove environment-specific customer literals from expected-result cells."""

from pathlib import Path

from openpyxl import load_workbook

workbook_path = Path(__file__).resolve().parents[1] / "test_case.xlsx"
workbook = load_workbook(workbook_path)
sheet = workbook["自动化测试用例"]
headers: dict[str, int] = {}
for cell in sheet[1]:
    if isinstance(cell.value, str) and isinstance(cell.column, int):
        headers[cell.value] = cell.column

replacements = {
    "TC-HOME-004": "可搜索到对应顾客：${SEEDED_CUSTOMER_NAME}",
    "TC-HOME-005": "可搜索到对应顾客：${SEEDED_CUSTOMER_NAME}",
    "TC-HOME-006": "结果手机号包含${SEEDED_CUSTOMER_PHONE_PART}",
    "TC-HOME-007": "结果手机号等于${SEEDED_CUSTOMER_PHONE}",
}

for row in range(2, sheet.max_row + 1):
    case_id = str(sheet.cell(row, headers["用例ID"]).value or "")
    if case_id in replacements:
        sheet.cell(row, headers["期望结果"], replacements[case_id])

workbook.save(workbook_path)
print("Sanitized environment-specific customer expectations in test_case.xlsx.")
